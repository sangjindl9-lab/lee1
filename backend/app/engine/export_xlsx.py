"""한국식 물량산출조서 Excel (10개 시트). A4/A3 가로, 머리글 반복, 페이지번호, 합계 수식."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
FONT = Font(name="맑은 고딕", size=10)
BOLD = Font(name="맑은 고딕", size=10, bold=True)
TITLE = Font(name="맑은 고딕", size=16, bold=True)
STATUS_KO = {"auto": "자동 검출", "verified": "사용자 검증 완료", "edited": "사용자 수정 완료", "review": "추가 검토 필요",
             "insufficient": "도면 정보 부족", "impossible": "산출 불가", "approved": "승인", "unresolved": "미확인", "excluded": "제외", "manual": "수동 추가"}
INSTALL_KO = {"new": "신설", "existing": "기존", "remove": "철거", "relocate": "이설"}


def _header_block(ws, title: str, meta: dict, ncols: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(1, 1, title).font = TITLE
    ws.cell(1, 1).alignment = Alignment(horizontal="center")
    info = [("공사명", meta.get("name", "")), ("발주처", meta.get("client", "")), ("작성자", meta.get("author", "")),
            ("검토자", meta.get("reviewer", "")), ("작성일자", meta.get("date", "")), ("도면 개정번호", meta.get("rev", ""))]
    r = 2
    for i in range(0, len(info), 3):
        c = 1
        for k, v in info[i:i + 3]:
            ws.cell(r, c, k).font = BOLD
            ws.cell(r, c + 1, v).font = FONT
            c += max(2, ncols // 3)
        r += 1
    return r + 1


def _table(ws, start_row: int, headers: list[str], rows: list[list], widths: list[int] | None = None, sum_cols: list[int] | None = None):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(start_row, i, h)
        c.font, c.fill, c.border, c.alignment = BOLD, HEAD_FILL, BORDER, Alignment(horizontal="center", vertical="center", wrap_text=True)
    r = start_row + 1
    for row in rows:
        for i, v in enumerate(row, start=1):
            c = ws.cell(r, i, v)
            c.font, c.border = FONT, BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="right" if isinstance(v, (int, float)) else "left")
        r += 1
    if sum_cols and rows:
        ws.cell(r, 1, "합계").font = BOLD
        ws.cell(r, 1).border = BORDER
        for col in sum_cols:
            L = get_column_letter(col)
            c = ws.cell(r, col, f"=SUM({L}{start_row + 1}:{L}{r - 1})")
            c.font, c.border = BOLD, BORDER
        for i in range(1, len(headers) + 1):
            ws.cell(r, i).border = BORDER
        r += 1
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.print_title_rows = f"{start_row}:{start_row}"
    ws.oddFooter.center.text = "&P / &N"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    return r


def build_workbook(data: dict, out: Path, a3: bool = False) -> Path:
    """data: {meta, items, items_by_sheet, routes, detections, legend, issues}"""
    wb = Workbook()
    meta = data["meta"]
    items = data["items"]

    def item_rows(filter_fn=lambda q: True, with_disc=True):
        rows = []
        for i, q in enumerate([q for q in items if filter_fn(q)], start=1):
            rows.append([i, q["discipline"] or "", q.get("sub_discipline") or "", q["item_name"], q["spec"] or "", q["unit"] or "",
                         q["formula"] or "", q["qty_total"], ", ".join(q["drawing_nos"] or []), q["floor_zone"] or "",
                         f"검출 {len(q.get('evidence', []))}건 / {STATUS_KO.get(q['status'], q['status'])}", (q.get("note") or "") + (f" [{INSTALL_KO.get(q['install_type'],'')}]" if q.get("install_type") != "new" else "")])
        return rows

    H = ["번호", "공종", "세부공종", "품명", "규격", "단위", "산출식", "산출수량", "도면번호", "층·구간", "산출근거", "비고"]
    W = [5, 10, 10, 24, 22, 6, 40, 10, 14, 10, 20, 18]

    # 1. 공사 전체 물량집계표
    ws = wb.active
    ws.title = "1.전체집계표"
    r = _header_block(ws, "공사 전체 물량집계표", meta, 8)
    agg: dict[tuple, dict] = {}
    for q in items:
        k = (q["discipline"] or "", q["item_name"], q["spec"] or "", q["unit"] or "", q.get("install_type", "new"))
        a = agg.setdefault(k, {"qty": 0.0, "pending": False})
        if q["qty_total"] is None:
            a["pending"] = True
        else:
            a["qty"] += q["qty_total"]
    rows = [[i, k[0], k[1], k[2], k[3], (a["qty"] if not a["pending"] else None), INSTALL_KO.get(k[4], ""), "미산출(검토 필요)" if a["pending"] else ""]
            for i, (k, a) in enumerate(sorted(agg.items()), start=1)]
    _table(ws, r, ["번호", "공종", "품명", "규격", "단위", "수량", "구분", "비고"], rows, [5, 12, 26, 24, 6, 12, 8, 20], sum_cols=[6])

    # 2/3. 전기·정보통신 집계
    for name, disc in (("2.전기집계표", "전기"), ("3.정보통신집계표", "정보통신")):
        ws = wb.create_sheet(name)
        r = _header_block(ws, f"{disc}공사 물량집계표", meta, 8)
        rows = [[i, k[1], k[2], k[3], (a["qty"] if not a["pending"] else None), INSTALL_KO.get(k[4], ""), "미산출" if a["pending"] else ""]
                for i, (k, a) in enumerate([kv for kv in sorted(agg.items()) if kv[0][0] == disc], start=1)]
        _table(ws, r, ["번호", "품명", "규격", "단위", "수량", "구분", "비고"], rows, [5, 26, 24, 6, 12, 8, 20], sum_cols=[5])

    # 4. 공종별 상세 수량산출서
    ws = wb.create_sheet("4.상세산출서")
    r = _header_block(ws, "공종별 상세 수량산출서", meta, len(H))
    _table(ws, r, H, item_rows(), W, sum_cols=[8])

    # 5. 도면별 물량산출서
    ws = wb.create_sheet("5.도면별산출서")
    r = _header_block(ws, "도면별 물량산출서", meta, 8)
    rows = [[i, d["drawing_no"], d["drawing_title"] or "", d["item_name"], d["spec"] or "", d["unit"] or "", d["qty"], d["basis"]]
            for i, d in enumerate(data["items_by_sheet"], start=1)]
    _table(ws, r, ["번호", "도면번호", "도면명", "품명", "규격", "단위", "수량", "근거"], rows, [5, 12, 24, 24, 20, 6, 10, 30], sum_cols=[7])

    # 6. 배관·케이블 길이 산출표
    ws = wb.create_sheet("6.길이산출표")
    r = _header_block(ws, "배관·트레이·케이블 길이 산출표", meta, 11)
    rows = [[i, rt["drawing_no"], rt["kind_ko"], rt["item_name"], rt["spec"] or "", rt["segment_count"], rt["horizontal_m"], rt["vertical_m"],
             rt["slack_m"], rt["circuits"], STATUS_KO.get(rt["status"], rt["status"])] for i, rt in enumerate(data["routes"], start=1)]
    _table(ws, r, ["번호", "도면번호", "구분", "품명", "규격", "구간수", "도면길이(m)", "수직길이(m)", "여장(m)", "회선", "상태"], rows,
           [5, 12, 8, 24, 20, 7, 12, 12, 10, 16, 14], sum_cols=[7, 8, 9])

    # 7. 장비 및 자재 집계표
    ws = wb.create_sheet("7.장비자재집계")
    r = _header_block(ws, "장비 및 자재 집계표", meta, 7)
    rows = [[i, k[1], k[2], k[3], (a["qty"] if not a["pending"] else None), "관급" if k[1] in data.get("gov_items", ()) else "사급", INSTALL_KO.get(k[4], "")]
            for i, (k, a) in enumerate(sorted(agg.items()), start=1)]
    _table(ws, r, ["번호", "품명", "규격", "단위", "수량", "관급/사급", "구분"], rows, [5, 26, 24, 6, 12, 10, 8], sum_cols=[5])

    # 8. 기호 검출 목록
    ws = wb.create_sheet("8.기호검출목록")
    r = _header_block(ws, "기호 검출 목록 (근거)", meta, 10)
    rows = [[d["id"], d["drawing_no"], d["item_name"], d["spec"] or "", d["seq_no"], round(d["x"], 2), round(d["y"], 2), d["method"], d["confidence"],
             STATUS_KO.get(d["status"], d["status"])] for d in data["detections"]]
    _table(ws, r, ["검출ID", "도면번호", "품명", "규격", "번호", "X", "Y", "방법", "신뢰도", "상태"], rows, [8, 12, 24, 20, 6, 12, 12, 10, 8, 14])

    # 9. 미확정 및 오류 검토 목록
    ws = wb.create_sheet("9.검토목록")
    r = _header_block(ws, "미확정 및 오류 검토 목록", meta, 5)
    rows = [[i, iss["severity"], iss["code"], iss["message"], "해결" if iss["resolved"] else "미해결"] for i, iss in enumerate(data["issues"], start=1)]
    rows += [[len(rows) + i, "warn", "QTY_" + q["status"].upper(), f"{q['item_name']} {q['spec'] or ''}: {STATUS_KO.get(q['status'])} — {q.get('note') or q['formula']}", "미해결"]
             for i, q in enumerate([q for q in items if q["status"] in ("review", "insufficient", "impossible")], start=1)]
    _table(ws, r, ["번호", "심각도", "코드", "내용", "상태"], rows, [5, 8, 20, 90, 8])

    # 10. 범례 매칭표
    ws = wb.create_sheet("10.범례매칭표")
    r = _header_block(ws, "범례 매칭표", meta, 10)
    rows = [[i, L["table_title"], L["symbol_name"] or "", L["block_name"] or "", L["item_name"] or "", L["spec"] or "", L["unit"] or "",
             L["category"], L["confidence"], STATUS_KO.get(L["status"], L["status"]), L["detected_count"]] for i, L in enumerate(data["legend"], start=1)]
    _table(ws, r, ["번호", "범례표", "기호명", "CAD블록", "품명", "규격", "단위", "범주", "신뢰도", "상태", "검출수"], rows, [5, 18, 10, 16, 24, 20, 6, 8, 8, 12, 8])

    if a3:
        for w in wb.worksheets:
            w.page_setup.paperSize = w.PAPERSIZE_A3
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def fill_company_template(template: Path, mapping: dict[str, str], items: list[dict], out: Path, sheet_name: str | None = None, start_row: int | None = None) -> Path:
    """회사 표준 양식에 컬럼 매핑 출력. mapping: {"A":"item_name","B":"spec",...}. start_row 미지정 시 첫 빈 행."""
    wb = load_workbook(template)
    ws = wb[sheet_name] if sheet_name else wb.active
    r = start_row or (ws.max_row + 1)
    for q in items:
        for col, field in mapping.items():
            v = q.get(field)
            if field == "drawing_nos" and isinstance(v, list):
                v = ", ".join(v)
            ws[f"{col}{r}"] = v
        r += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def build_estimate_workbook(meta: dict, est: dict, out: Path) -> Path:
    """원가계산서(총괄) / 일위대가(품목별 재료비·노무비·근거) / 관급·사급 자재 집계 / 미산정 목록."""
    wb = Workbook()
    ws = wb.active
    ws.title = "1.원가계산서"
    r = _header_block(ws, "공사원가계산서", meta, 4)
    rows = [["직접재료비(사급)", est.get("direct_material"), "", "품목별 자재단가 × 수량"],
            ["관급자재(별도)", est.get("gov_material"), "", "관급 품목 합계(원가 제외)"],
            ["직접노무비", est.get("direct_labor"), "", "표준품셈 노무량 × 노임단가 × 수량"]]
    for name, o in (est.get("overhead") or {}).items():
        rows.append([name, o.get("amount"), f"{o.get('pct', '')}% / {o.get('base', '')}" if o.get("amount") is not None else "", o.get("source") or o.get("note") or ""])
    rows += [["공급가액", est.get("supply"), "", "직접공사비 + 제경비"], [f"부가가치세({est.get('vat_pct', 10)}%)", est.get("vat"), "", ""], ["총공사비", est.get("total"), "", ""]]
    _table(ws, r, ["비목", "금액(원)", "비율/기준", "적용 근거"], [[a, (None if b is None else round(b)), c, d] for a, b, c, d in rows], [22, 16, 20, 50])

    ws = wb.create_sheet("2.일위대가")
    r = _header_block(ws, "일위대가 및 품목별 원가", meta, 11)
    rows = []
    for i, L in enumerate(est.get("lines") or [], start=1):
        rows.append([i, L["item_name"], L.get("spec") or "", L.get("unit") or "", L.get("qty"), L.get("material_unit"), L.get("material_cost"),
                     L.get("labor_unit"), L.get("labor_cost"), L.get("norm") or "", "; ".join(L.get("notes") or []) + (f" / 자재: {L.get('material_source')}" if L.get("material_source") else "")])
    _table(ws, r, ["번호", "품명", "규격", "단위", "수량", "재료비 단가", "재료비", "노무비 단가", "노무비", "적용 품셈", "비고/근거"], rows,
           [5, 24, 20, 6, 10, 12, 14, 12, 14, 24, 40], sum_cols=[7, 9])

    ws = wb.create_sheet("3.관급사급집계")
    r = _header_block(ws, "관급·사급 자재 집계", meta, 6)
    rows = [[i, L["item_name"], L.get("spec") or "", L.get("unit") or "", L.get("qty"), "관급" if "관급자재" in (L.get("notes") or []) else "사급"]
            for i, L in enumerate(est.get("lines") or [], start=1) if L.get("material_cost") is not None]
    _table(ws, r, ["번호", "품명", "규격", "단위", "수량", "구분"], rows, [5, 24, 20, 6, 10, 8])

    ws = wb.create_sheet("4.미산정목록")
    r = _header_block(ws, "미산정·검토 필요 항목", meta, 3)
    rows = [[i, L["item_name"] + " " + (L.get("spec") or ""), "; ".join(L.get("notes") or [])] for i, L in enumerate(est.get("lines") or [], start=1)
            if L.get("material_cost") is None or L.get("labor_cost") is None]
    _table(ws, r, ["번호", "품목", "사유"], rows, [5, 40, 70])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
