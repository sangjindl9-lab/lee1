"""계약내역 버전 / 품목 조회·편집 / 원가 비목 / 일위대가."""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session

from ..auth import audit, get_current_user, get_project_for
from ..db import get_db
from ..config import settings
from ..models import ContractVersion, CostRateRule, DesignChange, EstimateItem, PriceTable, ProgressClaim, SourceFile, UnitPriceAnalysis, User, WorkbookSheet
from ..schemas import CostRuleIn, FitContractIn, ItemAdd, ItemEdit, PriceAdjustIn
from ..services.rate_rows import evaluate_base
from ..services.costcalc import is_always_fixed
from ..services.calc import D, line_amounts, project_rounding, s

router = APIRouter(prefix="/api")


def item_out(it: EstimateItem) -> dict:
    return {
        "id": it.id, "uid": it.uid, "seq": it.seq, "row_type": it.row_type, "level": it.level, "section_path": it.section_path,
        "sheet": it.sheet, "source_row": it.source_row, "work_code": it.work_code, "item_code": it.item_code, "code": it.code,
        "name": it.name, "spec": it.spec, "unit": it.unit, "qty": s(it.qty),
        "mat_price": s(it.mat_price), "mat_amt": s(it.mat_amt), "lab_price": s(it.lab_price), "lab_amt": s(it.lab_amt),
        "exp_price": s(it.exp_price), "exp_amt": s(it.exp_amt), "tot_price": s(it.tot_price), "tot_amt": s(it.tot_amt),
        "note": it.note, "unit_price_ref": it.unit_price_ref, "price_source": it.price_source, "cost_region": it.cost_region,
        "is_new": it.is_new, "is_deleted": it.is_deleted, "issues": it.issues,
    }


def version_out(v: ContractVersion) -> dict:
    return {"id": v.id, "project_id": v.project_id, "version_no": v.version_no, "label": v.label, "status": v.status,
            "source_file_id": v.source_file_id, "design_change_id": v.design_change_id, "totals": v.totals,
            "approved_at": v.approved_at.isoformat() if v.approved_at else None, "created_at": v.created_at.isoformat(),
            "item_count": sum(1 for i in v.items if i.row_type == "item")}


def _version(db, vid, user, need="view") -> ContractVersion:
    v = db.get(ContractVersion, vid)
    if not v:
        raise HTTPException(404, "계약내역 버전을 찾을 수 없습니다.")
    get_project_for(db, v.project_id, user, need)
    return v


@router.get("/projects/{pid}/versions")
def list_versions(pid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    get_project_for(db, pid, user)
    return [version_out(v) for v in db.query(ContractVersion).filter_by(project_id=pid).order_by(ContractVersion.version_no)]


@router.get("/versions/{vid}")
def get_version(vid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return version_out(_version(db, vid, user))


@router.get("/versions/{vid}/items")
def version_items(vid: int, q: str | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    v = _version(db, vid, user)
    items = v.items
    if q:
        ql = q.lower()
        items = [i for i in items if i.row_type == "section" or ql in (i.name or "").lower() or ql in (i.spec or "").lower()]
    return {"version": version_out(v), "items": [item_out(i) for i in items]}


def _editable(db: Session, v: ContractVersion):
    """편집 가능 여부: 확정 기성이나 승인된 설계변경이 이 버전을 참조하면 직접 수정 금지(설계변경으로)."""
    if v.status == "draft":
        return
    if db.query(ProgressClaim).filter_by(contract_version_id=v.id, status="confirmed").count():
        raise HTTPException(400, "이 계약내역으로 확정된 기성 회차가 있어 직접 수정할 수 없습니다. 설계변경으로 처리하세요.")
    if db.query(DesignChange).filter_by(base_version_id=v.id, status="approved").count():
        raise HTTPException(400, "이 계약내역을 기준으로 승인된 설계변경이 있어 직접 수정할 수 없습니다. 최신 버전을 편집하거나 설계변경으로 처리하세요.")


@router.get("/versions/{vid}/editable")
def version_editable(vid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    v = _version(db, vid, user)
    try:
        _editable(db, v)
        return {"editable": True, "reason": ""}
    except HTTPException as e:
        return {"editable": False, "reason": e.detail}


@router.put("/versions/{vid}/items/{uid}")
def edit_item(vid: int, uid: str, body: ItemEdit, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """품목 수정 (수량·단가·명칭). 금액은 현장 반올림 규칙으로 재계산. 확정 버전은 기성/설계변경 참조가 없을 때만."""
    v = _version(db, vid, user, "edit")
    _editable(db, v)
    p = get_project_for(db, v.project_id, user, "edit")
    it = db.query(EstimateItem).filter_by(version_id=vid, uid=uid).first()
    if not it:
        raise HTTPException(404, "품목이 없습니다.")
    before = item_out(it)
    for k in ("name", "spec", "unit", "note"):
        val = getattr(body, k)
        if val is not None:
            setattr(it, k, val)
    for k in ("qty", "mat_price", "lab_price", "exp_price"):
        val = getattr(body, k)
        if val is not None:
            setattr(it, k, Decimal(str(val)))
    if it.row_type == "item":
        am = line_amounts(D(it.qty), it.mat_price, it.lab_price, it.exp_price, project_rounding(p))
        it.mat_amt, it.lab_amt, it.exp_amt = am["mat_amt"], am["lab_amt"], am["exp_amt"]
        it.tot_price = D(it.mat_price) + D(it.lab_price) + D(it.exp_price)
        it.tot_amt = am["tot_amt"]
    _recalc_totals(v)
    audit(db, user, "item.edit", v.project_id, target=f"v{v.version_no}/{it.name}", before=before, after=item_out(it), reason=body.reason)
    db.commit()
    return item_out(it)


@router.post("/versions/{vid}/items")
def add_item(vid: int, body: ItemAdd, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """품목 추가: after_uid 다음(또는 해당 공종 끝)에 삽입."""
    import hashlib
    v = _version(db, vid, user, "edit")
    _editable(db, v)
    p = get_project_for(db, v.project_id, user, "edit")
    items = sorted(v.items, key=lambda i: i.seq)
    anchor = next((i for i in items if i.uid == body.after_uid), None) if body.after_uid else None
    if anchor is None:
        # 공종 경로의 마지막 품목 뒤
        sec = [i for i in items if i.section_path == body.section_path]
        anchor = sec[-1] if sec else None
    if anchor is None:
        raise HTTPException(400, "삽입 위치(공종)를 찾지 못했습니다.")
    path = anchor.section_path if anchor.row_type == "item" else (body.section_path or anchor.section_path)
    pos = anchor.seq
    for i in items:
        if i.seq > pos:
            i.seq += 1
    uid = hashlib.sha1(f"manual|{vid}|{pos}|{body.name}|{body.spec}|{datetime.utcnow().isoformat()}".encode()).hexdigest()[:16]
    am = line_amounts(D(body.qty), body.mat_price, body.lab_price, body.exp_price, project_rounding(p))
    it = EstimateItem(version_id=vid, uid=uid, seq=pos + 1, row_type="item", level=len(path) + 1, section_path=path,
                      sheet=anchor.sheet, name=body.name, spec=body.spec, unit=body.unit, qty=D(body.qty),
                      mat_price=D(body.mat_price), lab_price=D(body.lab_price), exp_price=D(body.exp_price),
                      mat_amt=am["mat_amt"], lab_amt=am["lab_amt"], exp_amt=am["exp_amt"],
                      tot_price=D(body.mat_price) + D(body.lab_price) + D(body.exp_price), tot_amt=am["tot_amt"], note=body.note,
                      price_source={"manual": True})
    v.items.append(it)
    _recalc_totals(v)
    audit(db, user, "item.add", v.project_id, target=f"v{v.version_no}/{it.name}", after=item_out(it), reason=body.reason)
    db.commit()
    return item_out(it)


@router.delete("/versions/{vid}/items/{uid}")
def delete_item(vid: int, uid: str, reason: str = "", db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    v = _version(db, vid, user, "edit")
    _editable(db, v)
    it = db.query(EstimateItem).filter_by(version_id=vid, uid=uid).first()
    if not it:
        raise HTTPException(404, "품목이 없습니다.")
    before = item_out(it)
    if it.price_source.get("manual") or v.status == "draft" and not it.source_row:
        v.items.remove(it)
        db.delete(it)          # 수동 추가 행은 완전 삭제
    else:
        it.is_deleted = True   # 원본 행은 삭제 표시 (원본 양식 출력 시 수량 0 처리)
    _recalc_totals(v)
    audit(db, user, "item.delete", v.project_id, target=f"v{v.version_no}/{before['name']}", before=before, reason=reason)
    db.commit()
    return {"ok": True, "totals": v.totals}


@router.get("/versions/{vid}/export")
def export_version(vid: int, delete_mode: str = "zero", db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """계약내역 버전을 업로드한 원본 엑셀 양식 그대로 출력 (수정·추가·삭제된 셀만 반영, 나머지 시트·서식·수식 유지)."""
    from pathlib import Path
    from ..excel.analyzer import analyze_file
    from ..excel.export_version import write_version_original
    v = _version(db, vid, user)
    p = get_project_for(db, v.project_id, user)
    sf = db.get(SourceFile, v.source_file_id) if v.source_file_id else None
    if not sf:
        raise HTTPException(400, "이 버전은 원본 파일이 없어 원본 양식으로 출력할 수 없습니다.")
    sheet_maps = {ws.name: ws.column_map for ws in db.query(WorkbookSheet).filter_by(file_id=sf.id) if ws.column_map}
    # 원본 값(가져오기 시점) 재추출 → 바뀐 셀만 기록
    role_over = {ws.name: ws.user_role for ws in sf.sheets if ws.user_role}
    map_over = {ws.name: ws.user_column_map for ws in sf.sheets if ws.user_column_map}
    res = analyze_file(sf.stored_path, str(settings.storage_dir / "work"), role_over or None, map_over or None)
    original = {}
    for ps in res["_parsed"].values():
        for it in ps.items:
            original[it.uid] = {"name": it.name, "spec": it.spec, "unit": it.unit, "note": it.note, "qty": it.qty,
                                "mat_price": it.mat_price, "lab_price": it.lab_price, "exp_price": it.exp_price, "tot_price": it.tot_price}
    out_dir = settings.storage_dir / f"project_{p.id}" / "exports"
    out_dir.mkdir(parents=True, exist_ok=True)
    ext = Path(sf.stored_path).suffix.lower()
    stem = Path(sf.original_name).stem
    out = out_dir / f"{stem}_v{v.version_no}_{v.label}_{datetime.now().strftime('%Y%m%d_%H%M')}{ext}"
    cost_over = [{"sheet": r.sheet, "row": r.source_row, "col": r.amount_col, "amount": s(r.amount), "name": r.name,
                  "force": bool(r.fixed and r.name.replace(" ", "").startswith("이윤"))}
                 for r in db.query(CostRateRule).filter_by(version_id=v.id) if r.kind == "line" and r.amount is not None and r.source_row]
    stats = write_version_original(Path(sf.stored_path), out, v, sheet_maps, original, delete_mode,
                                   title_suffix=f"v{v.version_no} {v.label}" if v.version_no else "", cost_overrides=cost_over)
    audit(db, user, "version.export", p.id, target=f"v{v.version_no}", after={"file": out.name, **{k: stats[k] for k in ("modified", "added", "deleted")}})
    db.commit()
    media = "application/vnd.ms-excel.sheet.macroEnabled.12" if ext == ".xlsm" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return FileResponse(out, filename=out.name, media_type=media, headers={"X-Export-Stats": f"modified={stats['modified']};added={stats['added']};deleted={stats['deleted']}"})


def _restore_rate(amt: Decimal, base: Decimal, rule) -> Decimal:
    """금액/기준금액에서 원래 요율(2%, 1.5%, 0.25% …)을 복원: 자릿수가 적은 비율부터 금액이 재현되는 것을 택한다."""
    from ..services.calc import round_by
    raw = amt / base * 100
    for d in range(0, 5):
        cand = raw.quantize(Decimal(1).scaleb(-d))
        if cand > 0 and round_by(base * cand / 100, rule) == amt:
            return cand / 100
    return amt / base


def _rate_lookup(v: ContractVersion, new_vals: dict):
    """원본 통합문서(캐시값)를 읽어 (행, 열) -> 값 조회 함수를 시트별로 만든다. 금액 열은 조정된 금액으로 대체."""
    import openpyxl
    from sqlalchemy.orm import object_session
    db = object_session(v)
    sf = db.get(SourceFile, v.source_file_id) if (db and v.source_file_id) else None
    if not sf or not Path(sf.stored_path).exists():
        return None
    try:
        wb = openpyxl.load_workbook(sf.stored_path, data_only=True, read_only=False)
    except Exception:
        return None
    maps = {ws.name: (ws.user_column_map or ws.column_map or {}) for ws in sf.sheets}
    cache: dict[str, dict] = {}

    def for_sheet(sheet: str):
        if sheet not in wb.sheetnames:
            return lambda r, c: None
        ws = wb[sheet]
        cm = maps.get(sheet, {})
        col_field = {}
        fields = cm.get("fields", cm)
        for f in ("mat_amt", "lab_amt", "exp_amt"):
            c = fields.get(f)
            if isinstance(c, dict):
                c = c.get("col")
            if isinstance(c, str) and c.isalpha():
                from openpyxl.utils import column_index_from_string
                c = column_index_from_string(c)
            if c:
                col_field[int(c)] = f

        def lookup(r: int, c: int):
            f = col_field.get(c)
            if f:
                nv = new_vals.get((sheet, r))
                if nv is not None:
                    return nv[f]
            return ws.cell(r, c).value
        return lookup
    return for_sheet


def _apply_pct(v: ContractVersion, rounding: dict, body, apply: bool) -> dict:
    """단가 % 조정 공통 로직. apply=False 면 미리보기."""
    from ..services.calc import round_by
    prule = {"func": body.price_round, "digits": body.price_digits}
    pcts = {"mat_price": Decimal(str(body.mat_pct)) / 100, "lab_price": Decimal(str(body.lab_pct)) / 100, "exp_price": Decimal(str(body.exp_pct)) / 100}
    before_tot = {k: Decimal(0) for k in ("mat_amt", "lab_amt", "exp_amt", "tot_amt")}
    after_tot = {k: Decimal(0) for k in before_tot}
    n = 0
    skipped = []
    rate_rows = []
    # 일반 행 먼저 처리 → 요율행은 조정된 금액으로 기준금액 수식을 다시 평가
    new_vals: dict[tuple[str, int], dict[str, Decimal]] = {}   # (시트, 행) -> 조정 후 금액
    for it in v.items:
        if it.row_type != "item" or it.is_deleted:
            continue
        if body.section_prefix and it.section_path[:len(body.section_prefix)] != body.section_prefix:
            continue
        # 금액이 수량×단가와 다른 행(공구손료 등 요율 산정 행): 기준 비목의 적용률로 비례 조정
        cur_am = line_amounts(D(it.qty), it.mat_price, it.lab_price, it.exp_price, rounding)
        if any(abs(cur_am[k] - D(getattr(it, k))) > 1 for k in ("mat_amt", "lab_amt", "exp_amt")):
            rate_rows.append(it)
            continue
        for k in before_tot:
            before_tot[k] += D(getattr(it, k))
        new = {}
        for f, pct in pcts.items():
            cur = getattr(it, f)
            new[f] = round_by(D(cur) * pct, prule) if (cur is not None and pct != 1) else cur
        am = line_amounts(D(it.qty), new["mat_price"], new["lab_price"], new["exp_price"], rounding)
        for k in after_tot:
            after_tot[k] += am[k]
        if it.sheet and it.source_row:
            new_vals[(it.sheet, int(it.source_row))] = am
        if not body.only_preview:
            changed = any(D(new[f]) != D(getattr(it, f)) for f in pcts)
            if changed:
                n += 1
                it.mat_price, it.lab_price, it.exp_price = new["mat_price"], new["lab_price"], new["exp_price"]
                it.mat_amt, it.lab_amt, it.exp_amt, it.tot_amt = am["mat_amt"], am["lab_amt"], am["exp_amt"], am["tot_amt"]
                it.tot_price = D(it.mat_price) + D(it.lab_price) + D(it.exp_price)
                src = dict(it.price_source or {})
                src["adjusted"] = {"mat_pct": body.mat_pct, "lab_pct": body.lab_pct, "exp_pct": body.exp_pct}
                it.price_source = src
        else:
            n += 1
    # 요율행(공구손료 "노무비의 2%" 등): 단가칸 수식(=SUMIF(...노무비 열...))을 조정된 금액으로 다시 평가해 기준금액을 구하고
    # 금액 = 기준금액 × 요율. 수식을 평가하지 못하면 기준 비목의 적용률로 비례 조정. 단가 셀의 수식은 출력 시 그대로 둔다(rate_row).
    lookup = _rate_lookup(v, new_vals) if rate_rows else None
    for it in rate_rows:
        text = f"{it.name} {it.spec}".replace(" ", "")
        dep = "lab" if any(w in text for w in ("인력품", "노무", "품의", "공구손료")) else "mat"
        pct = pcts["lab_price"] if dep == "lab" else pcts["mat_price"]
        for k in before_tot:
            before_tot[k] += D(getattr(it, k))
        new_am = {}
        new_price = {}
        for k, pf in (("mat_amt", "mat_price"), ("lab_amt", "lab_price"), ("exp_amt", "exp_price")):
            old_amt, old_price = D(getattr(it, k)), D(getattr(it, pf))
            if old_amt == 0:
                new_am[k], new_price[pf] = old_amt, getattr(it, pf)
                continue
            base_a = None
            formula = (it.price_source or {}).get(pf)
            if lookup and isinstance(formula, str) and formula.startswith("=") and it.sheet:
                base_a = evaluate_base(formula, lookup(it.sheet))
            if base_a is not None and old_price and D(it.qty) == 1:
                rate = _restore_rate(old_amt, old_price, rounding.get("amount"))   # 2%, 1.5% 등 요율 복원
                new_price[pf] = base_a
                new_am[k] = round_by(base_a * rate, rounding.get("amount"))
            else:
                new_price[pf] = getattr(it, pf)
                new_am[k] = round_by(old_amt * pct, rounding.get("amount"))
        new_am["tot_amt"] = new_am["mat_amt"] + new_am["lab_amt"] + new_am["exp_amt"]
        for k in after_tot:
            after_tot[k] += new_am[k]
        if it.sheet and it.source_row:
            new_vals[(it.sheet, int(it.source_row))] = new_am
        if apply and pct != 1:
            it.mat_amt, it.lab_amt, it.exp_amt, it.tot_amt = new_am["mat_amt"], new_am["lab_amt"], new_am["exp_amt"], new_am["tot_amt"]
            it.mat_price, it.lab_price, it.exp_price = new_price["mat_price"], new_price["lab_price"], new_price["exp_price"]
            it.tot_price = D(it.mat_price) + D(it.lab_price) + D(it.exp_price)
            src = dict(it.price_source or {})
            src["rate_row"] = True
            it.price_source = src
            n += 1
        skipped.append(it.name)
    # 범위(section_prefix) 밖 행의 금액은 그대로 합계에 포함
    skipped_tot = {k: Decimal(0) for k in before_tot}
    for it in v.items:
        if it.row_type == "item" and not it.is_deleted and body.section_prefix and it.section_path[:len(body.section_prefix)] != body.section_prefix:
            for k in skipped_tot:
                skipped_tot[k] += D(getattr(it, k))
    if apply:
        _recalc_totals(v)
    return {"items": n, "skipped": skipped[:50], "skipped_count": len(skipped),
            "before": {k: s(x) for k, x in before_tot.items()}, "after": {k: s(x) for k, x in after_tot.items()},
            "after_all": {k: s(after_tot[k] + skipped_tot[k]) for k in after_tot},
            "diff": {k: s(after_tot[k] - before_tot[k]) for k in before_tot}, "preview": not apply}


@router.post("/versions/{vid}/adjust-prices")
def adjust_prices(vid: int, body: PriceAdjustIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """재료비/노무비/경비 단가를 퍼센트로 일괄 조정 (낙찰률·노임 상승률 적용 등). only_preview 면 합계 변화만 계산."""
    v = _version(db, vid, user, "edit")
    _editable(db, v)
    p = get_project_for(db, v.project_id, user, "edit")
    res = _apply_pct(v, project_rounding(p), body, apply=not body.only_preview)
    if not body.only_preview:
        audit(db, user, "items.adjust_prices", v.project_id, target=f"v{v.version_no}", before=res["before"], after=res["after"],
              reason=f"{body.reason} (재료 {body.mat_pct}% / 노무 {body.lab_pct}% / 경비 {body.exp_pct}%)")
        db.commit()
    return res


@router.post("/versions/{vid}/fit-contract")
def fit_contract(vid: int, body: FitContractIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """낙찰금액 맞추기: 재료/노무/경비 단가에 적용률을 곱한 뒤, 제비율을 재계산하고 남은 차액을 이윤으로 맞춘다.
    미리보기(only_preview=True)로 이윤·이윤율을 먼저 확인하고 적용한다."""
    from ..services.costcalc import fit_profit
    v = _version(db, vid, user, "edit")
    _editable(db, v)
    p = get_project_for(db, v.project_id, user, "edit")
    rounding = project_rounding(p)
    rules = db.query(CostRateRule).filter_by(project_id=p.id, version_id=v.id).order_by(CostRateRule.seq).all()
    if not rules or not any(r.name.replace(" ", "").startswith("이윤") for r in rules):
        raise HTTPException(400, "원가계산서에 '이윤' 비목이 없어 낙찰금액 맞추기를 할 수 없습니다. 원가계산서·제비율에서 비목을 확인하세요.")
    adj = _apply_pct(v, rounding, body, apply=not body.only_preview)
    after = adj["after_all"]
    # 기계경비는 원가계산서 금액(외부 산출)을 우선하고 경비 적용률을 곱한다
    exp_rule = next((r for r in rules if r.name.replace(" ", "").startswith("기계경비") and r.amount is not None), None)
    from ..services.calc import round_by
    exp_base = round_by(D(exp_rule.amount) * Decimal(str(body.exp_pct)) / 100, rounding.get("amount")) if exp_rule else D(after["exp_amt"])
    direct = {"mat": D(after["mat_amt"]), "lab": D(after["lab_amt"]), "exp": exp_base}
    # 조정 전(원본) 직접비 = 조정 범위 밖 합계 + 조정 전 합계
    orig = {k: D(after[k]) - D(adj["after"][k]) + D(adj["before"][k]) for k in ("mat_amt", "lab_amt")}
    direct_orig = {"mat": orig["mat_amt"], "lab": orig["lab_amt"], "exp": D(exp_rule.amount) if exp_rule else D(after["exp_amt"]) - D(adj["after"]["exp_amt"]) + D(adj["before"]["exp_amt"])}
    # 고정은 정산 비목(국민연금·노인장기요양·퇴직공제·안전관리비)만, 이윤·그 외 요율 비목은 다시 계산
    from ..services.costcalc import classify_line
    for r in rules:
        if r.name.replace(" ", "").startswith("이윤"):
            r.fixed = False
        elif r.kind == "line" and r.rate is not None and classify_line(r.name) == "line":
            r.fixed = is_always_fixed(r.name)
    fit = fit_profit(rules, direct, Decimal(str(body.target_amount)), body.include_vat, rounding.get("amount"), direct_orig)
    result = {"adjust": adj, "cost_sheet": fit["lines"], "summary": {k: s(x) if isinstance(x, Decimal) else x for k, x in fit["summary"].items()},
              "profit": s(fit["profit"]), "profit_rate": s(fit["profit_rate"]) if fit["profit_rate"] is not None else None,
              "profit_cap": s(fit["profit_cap"]) if fit["profit_cap"] is not None else None,
              "target": s(fit["target"]), "target_cost": s(fit["target_cost"]), "diff": s(fit["diff"]), "warnings": fit["warnings"], "preview": body.only_preview}
    if body.only_preview:
        db.rollback()
        return result
    if fit["profit"] < 0:
        db.rollback()
        raise HTTPException(400, "이윤이 음수가 되어 적용할 수 없습니다. 재료비/노무비 적용률을 낮추세요.")
    # 비목 금액 저장 + 이윤 고정
    by_id = {r.id: r for r in rules}
    for line in fit["lines"]:
        r = by_id.get(line["id"])
        if r is None:
            continue
        if line["amount"] is not None:
            r.amount = D(line["amount"])
        if line["cat"] == "profit":
            r.fixed = True
            r.source = f"낙찰금액 {s(Decimal(str(body.target_amount)))} 맞추기 (이윤 역산, 이윤율 {result['profit_rate']}%)"
    p.contract_amount = Decimal(str(body.target_amount))
    st = dict(p.settings or {})
    st["award"] = {"target_amount": body.target_amount, "include_vat": body.include_vat, "mat_pct": body.mat_pct, "lab_pct": body.lab_pct,
                   "exp_pct": body.exp_pct, "profit": s(fit["profit"]), "profit_rate": result["profit_rate"], "version_id": v.id}
    p.settings = st
    audit(db, user, "version.fit_contract", p.id, target=f"v{v.version_no}", before=adj["before"],
          after={"target": body.target_amount, "profit": s(fit["profit"]), "profit_rate": result["profit_rate"], "grand": result["summary"].get("grand")},
          reason=f"{body.reason} (재료 {body.mat_pct}% / 노무 {body.lab_pct}% / 경비 {body.exp_pct}%)")
    db.commit()
    return result


@router.post("/versions/{vid}/items/{uid}/restore")
def restore_item(vid: int, uid: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    v = _version(db, vid, user, "edit")
    _editable(db, v)
    it = db.query(EstimateItem).filter_by(version_id=vid, uid=uid).first()
    if not it:
        raise HTTPException(404)
    it.is_deleted = False
    _recalc_totals(v)
    audit(db, user, "item.restore", v.project_id, target=f"v{v.version_no}/{it.name}")
    db.commit()
    return item_out(it)


def _recalc_totals(v: ContractVersion):
    t = {k: Decimal(0) for k in ("mat_amt", "lab_amt", "exp_amt", "tot_amt")}
    for i in v.items:
        if i.row_type == "item" and not i.is_deleted:
            for k in t:
                t[k] += D(getattr(i, k))
    v.totals = {k: str(x) for k, x in t.items()}


@router.post("/versions/{vid}/approve")
def approve_version(vid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    v = _version(db, vid, user, "approve")
    v.status = "approved"
    v.approved_at = datetime.utcnow()
    audit(db, user, "version.approve", v.project_id, target=f"v{v.version_no} {v.label}")
    db.commit()
    return version_out(v)


@router.delete("/versions/{vid}")
def delete_version(vid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    v = _version(db, vid, user, "approve")
    delete_version_rows(db, v)
    audit(db, user, "version.delete", v.project_id, target=f"v{v.version_no} {v.label}")
    db.commit()
    return {"ok": True}


def delete_version_rows(db: Session, v: ContractVersion, others: set[int] | None = None) -> None:
    """계약내역 버전과 딸린 원가비목·설계변경을 삭제 (커밋은 호출측). others: 함께 지워지는 버전 id (서로 참조 허용)."""
    label = f"v{v.version_no} {v.label}"
    if db.query(ProgressClaim).filter_by(contract_version_id=v.id).count():
        raise HTTPException(400, f"계약내역 {label} 을(를) 사용하는 기성 회차가 있습니다. ④ 기성내역서에서 기성 회차를 먼저 삭제하세요.")
    others = others or set()
    dcs = db.query(DesignChange).filter((DesignChange.base_version_id == v.id) | (DesignChange.result_version_id == v.id)).all()
    if any(d.status == "approved" and d.result_version_id != v.id and d.result_version_id not in others for d in dcs):
        raise HTTPException(400, f"계약내역 {label} 을(를) 기준으로 승인된 설계변경이 있습니다. 변경 결과 버전을 먼저 삭제하세요.")
    for d in dcs:
        db.delete(d)   # 이 버전이 결과인 승인 설계변경 / 이 버전을 기준으로 한 작성중 설계변경은 함께 삭제
    db.query(CostRateRule).filter_by(version_id=v.id).delete()
    db.delete(v)


# ---------------------------------------------------------------- 원가 비목
def rule_out(r: CostRateRule) -> dict:
    return {"id": r.id, "version_id": r.version_id, "seq": r.seq, "name": r.name, "kind": r.kind, "rate": s(r.rate),
            "rate_base": r.rate_base, "base_text": r.base_text, "amount": s(r.amount), "fixed": r.fixed,
            "source": r.source, "effective_date": r.effective_date}


@router.get("/projects/{pid}/cost-rules")
def cost_rules(pid: int, version_id: int | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    get_project_for(db, pid, user)
    q = db.query(CostRateRule).filter_by(project_id=pid)
    if version_id:
        q = q.filter_by(version_id=version_id)
    return [rule_out(r) for r in q.order_by(CostRateRule.version_id, CostRateRule.seq)]


@router.put("/cost-rules/{rid}")
def update_rule(rid: int, body: CostRuleIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    r = db.get(CostRateRule, rid)
    if not r:
        raise HTTPException(404)
    get_project_for(db, r.project_id, user, "edit")
    before = rule_out(r)
    if body.rate is not None:
        r.rate = Decimal(str(body.rate))
    for k in ("rate_base", "fixed", "source", "effective_date"):
        if getattr(body, k) is not None:
            setattr(r, k, getattr(body, k))
    if body.amount is not None:
        r.amount = Decimal(str(body.amount))
    audit(db, user, "cost_rule.update", r.project_id, target=r.name, before=before, after=rule_out(r))
    db.commit()
    return rule_out(r)


@router.get("/projects/{pid}/unit-prices")
def unit_prices(pid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    get_project_for(db, pid, user)
    return [{"id": u.id, "sheet": u.sheet, "no": u.no, "title": u.title, "components": u.components, "total": u.total, "is_new": u.is_new}
            for u in db.query(UnitPriceAnalysis).filter_by(project_id=pid).order_by(UnitPriceAnalysis.id)]


@router.post("/projects/{pid}/unit-prices")
def create_unit_price(pid: int, body: dict, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """신규 일위대가 작성: components [{name,spec,unit,qty,mat_price,lab_price,exp_price}] → 합계 자동."""
    p = get_project_for(db, pid, user, "edit")
    comps = body.get("components", [])
    rounding = {"amount": {"func": "trunc", "digits": 1}}
    tot = {"mat_amt": Decimal(0), "lab_amt": Decimal(0), "exp_amt": Decimal(0)}
    for c in comps:
        am = line_amounts(D(c.get("qty")), c.get("mat_price"), c.get("lab_price"), c.get("exp_price"), rounding)
        for k in tot:
            tot[k] += am[k]
            c[k] = s(am[k])
    total = {k: s(v.quantize(Decimal(1))) for k, v in tot.items()}
    total["tot_amt"] = s(sum(tot.values()).quantize(Decimal(1)))
    u = UnitPriceAnalysis(project_id=p.id, sheet="신규", no=body.get("no") or f"신규 {db.query(UnitPriceAnalysis).filter_by(project_id=pid, is_new=True).count() + 1}호",
                          title=body.get("title", ""), components=comps, total=total, is_new=True)
    db.add(u)
    audit(db, user, "unit_price.create", p.id, target=u.no, after=total)
    db.commit()
    return {"id": u.id, "no": u.no, "title": u.title, "components": u.components, "total": u.total, "is_new": True}


# ---------------------------------------------------------------- 자재단가 / 노임 표
def _tokens_match(q: str, text: str) -> bool:
    t = text.lower().replace(" ", "")
    return all(tok in t for tok in q.lower().split() if tok)


def _row_text(r: dict) -> str:
    return " ".join(str(v) for k, v in r.items() if not k.startswith("_")) + " " + (r.get("_name") or "")
@router.get("/projects/{pid}/price-tables")
def price_tables(pid: int, kind: str | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    get_project_for(db, pid, user)
    q = db.query(PriceTable).filter_by(project_id=pid)
    if kind:
        q = q.filter_by(kind=kind)
    return [{"id": t.id, "sheet": t.sheet, "kind": t.kind, "title": t.title, "row_count": t.row_count, "source_file_id": t.source_file_id,
             "columns": t.columns} for t in q.order_by(PriceTable.kind, PriceTable.id)]


@router.get("/price-tables/{tid}")
def price_table(tid: int, q: str | None = None, limit: int = 3000, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    t = db.get(PriceTable, tid)
    if not t:
        raise HTTPException(404)
    get_project_for(db, t.project_id, user)
    rows = t.rows
    if q:
        rows = [r for r in rows if _tokens_match(q, _row_text(r))]
    return {"id": t.id, "sheet": t.sheet, "kind": t.kind, "title": t.title, "name_key": t.name_key, "columns": t.columns,
            "groups": t.groups, "rows": rows[:limit], "total": len(t.rows), "matched": len(rows)}


@router.get("/projects/{pid}/price-lookup")
def price_lookup(pid: int, q: str, limit: int = 30, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """품명/규격으로 자재단가·노임 표 통합 검색 (내역 편집·신규 품목 단가 참고용)."""
    get_project_for(db, pid, user)
    out = []
    for t in db.query(PriceTable).filter_by(project_id=pid).filter(PriceTable.kind.in_(["material", "labor"])):
        labels = {c["key"]: c["label"] for c in t.columns}
        for r in t.rows:
            if _tokens_match(q, _row_text(r)):
                nums = {labels[k]: v for k, v in r.items() if not k.startswith("_") and isinstance(v, (int, float)) and k in labels}
                out.append({"sheet": t.sheet, "kind": t.kind, "name": r.get("_name"), "row": r.get("_row"),
                            "text": {labels.get(k, k): v for k, v in r.items() if not k.startswith("_") and isinstance(v, str)}, "numbers": nums})
                if len(out) >= limit:
                    return out
    return out
