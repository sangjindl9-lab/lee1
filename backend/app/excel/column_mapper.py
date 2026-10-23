"""열 자동 인식 + 사용자 매핑 적용.

표준 필드:
  work_code  공종코드        item_code 품목코드       code 코드
  name 품명  spec 규격  unit 단위  qty 수량
  mat_price/mat_amt  재료비 단가/금액
  lab_qty/lab_price/lab_amt 노무비 (수량)/단가/금액
  exp_price/exp_amt  경비
  tot_price/tot_amt  합계 단가/금액
  price/amt          단일 단가/금액 (재료·노무 미구분 양식)
  note 비고
  orig_price/orig_qty/orig_amt, chg_*, diff_*  당초/변경/증감
"""
from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .loader import cell_text, norm_text

GROUPS = {
    "mat": ["재료비", "재료"],
    "lab": ["노무비", "노무"],
    "exp": ["경비"],
    "tot": ["계", "합계", "합계금액", "합계단가"],
    "orig": ["당초", "변경전"],
    "chg": ["변경", "변경후"],
    "diff": ["증감", "증·감", "증감(△)", "증△감"],
}
SUBS = {
    "price": ["단가"],
    "amt": ["금액"],
    "qty": ["수량"],
}
SINGLES = [
    ("work_code", ["공종코드"]),
    ("item_code", ["품목코드", "번호", "no"]),
    ("code", ["코드"]),
    ("name", ["품명", "명칭", "공종명", "품목", "공종"]),
    ("spec", ["규격"]),
    ("unit", ["단위"]),
    ("qty", ["수량"]),
    ("price", ["단가"]),
    ("amt", ["금액"]),
    ("note", ["비고"]),
    ("ref", ["일위대가", "호표", "산출근거"]),
]

FIELD_LABELS = {
    "work_code": "공종코드", "item_code": "품목코드", "code": "코드", "name": "품명", "spec": "규격", "unit": "단위",
    "qty": "수량", "mat_price": "재료비 단가", "mat_amt": "재료비 금액", "lab_qty": "노무 수량", "lab_price": "노무비 단가",
    "lab_amt": "노무비 금액", "exp_price": "경비 단가", "exp_amt": "경비 금액", "tot_price": "합계 단가", "tot_amt": "합계 금액",
    "price": "단가", "amt": "금액", "note": "비고", "ref": "참조",
    "orig_price": "당초 단가", "orig_qty": "당초 수량", "orig_amt": "당초 금액",
    "chg_price": "변경 단가", "chg_qty": "변경 수량", "chg_amt": "변경 금액",
    "diff_price": "증감 단가", "diff_qty": "증감 수량", "diff_amt": "증감 금액",
}


@dataclass
class ColumnMap:
    header_row: int
    header_rows: int                 # 헤더 행 수 (1 또는 2)
    data_start_row: int
    fields: dict[str, int] = field(default_factory=dict)   # field -> column index(1-based)
    raw_headers: dict[int, str] = field(default_factory=dict)  # col -> "재료비/단가"
    unmapped_cols: list[int] = field(default_factory=list)
    confidence: float = 0.0

    def col(self, f: str) -> int | None:
        return self.fields.get(f)

    def letter(self, f: str) -> str | None:
        c = self.fields.get(f)
        return get_column_letter(c) if c else None

    def has_cost_split(self) -> bool:
        return "mat_amt" in self.fields or "lab_amt" in self.fields

    def as_dict(self):
        return {
            "header_row": self.header_row,
            "header_rows": self.header_rows,
            "data_start_row": self.data_start_row,
            "confidence": round(self.confidence, 2),
            "fields": {k: {"col": v, "letter": get_column_letter(v), "label": FIELD_LABELS.get(k, k)}
                       for k, v in self.fields.items()},
            "raw_headers": {get_column_letter(k): v for k, v in self.raw_headers.items()},
            "unmapped": [get_column_letter(c) for c in self.unmapped_cols],
        }


def _merged_origin_value(ws: Worksheet, row: int, col: int):
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return ws.cell(row, col).value


def _match(text: str, words: list[str]) -> bool:
    t = norm_text(text)
    if not t:
        return False
    return any(norm_text(w) == t or (norm_text(w) in t and len(t) <= len(norm_text(w)) + 4) for w in words)


NAME_EXACT = {"품명", "명칭", "공종명", "품목", "품목명", "공종", "명칭및규격", "품명및규격", "비목", "공종명칭"}


def find_header_row(ws: Worksheet, hint: int | None = None, max_rows: int = 15) -> int | None:
    best, best_row = 0, None
    for r in range(1, min(ws.max_row, max_rows) + 1):
        score = 0
        has_name = False
        for c in range(1, min(ws.max_column, 40) + 1):
            t = cell_text(ws.cell(r, c).value)
            if not t or len(t) > 14:
                continue
            if norm_text(t) in NAME_EXACT:
                has_name = True
                score += 2
            for _, words in SINGLES:
                if _match(t, words):
                    score += 1
            for words in GROUPS.values():
                if _match(t, words):
                    score += 1
        if has_name:
            # 다음 행에 단가/금액/수량 보조 헤더가 있으면 2행 헤더의 첫 행 → 가산
            nxt = [norm_text(ws.cell(r + 1, c).value) for c in range(1, min(ws.max_column, 40) + 1)]
            if sum(1 for t in nxt if t in ("단가", "금액", "수량")) >= 2:
                score += 3
        if has_name and score > best:
            best, best_row = score, r
    return best_row or hint


def detect_columns(ws: Worksheet, header_row: int | None = None) -> ColumnMap | None:
    h = find_header_row(ws, header_row)
    if not h:
        return None
    max_col = min(ws.max_column, 60)
    # 2행 헤더 여부: 다음 행에 단가/금액/수량이 있으면 2행
    sub_row = h + 1
    sub_hits = sum(1 for c in range(1, max_col + 1)
                   if any(_match(cell_text(ws.cell(sub_row, c).value), w) for w in SUBS.values()))
    two_rows = sub_hits >= 2
    cm = ColumnMap(header_row=h, header_rows=2 if two_rows else 1, data_start_row=h + (2 if two_rows else 1))

    taken: set[str] = set()
    last_top = ""
    for c in range(1, max_col + 1):
        top = cell_text(_merged_origin_value(ws, h, c)) if two_rows else cell_text(ws.cell(h, c).value)
        sub = cell_text(ws.cell(sub_row, c).value) if two_rows else ""
        if two_rows:
            # 병합되지 않은 그룹 헤더(재료비 | 빈칸)도 오른쪽으로 이어진 것으로 본다
            if top:
                last_top = top
            elif sub and last_top and any(_match(last_top, w) for w in GROUPS.values()):
                top = last_top
            if not sub:
                last_top = top
        if not top and not sub:
            continue
        cm.raw_headers[c] = f"{top}/{sub}" if sub else top
        fld = None
        grp = next((g for g, words in GROUPS.items() if _match(top, words)), None)
        subk = next((s for s, words in SUBS.items() if _match(sub, words)), None) if sub else None
        if grp and subk:
            if grp == "tot" and subk == "qty":
                fld = None
            elif grp == "lab" and subk == "qty":
                fld = "lab_qty"
            else:
                fld = f"{grp}_{subk}"
        elif grp and not sub:
            # 단일 행에 '재료비'만 있으면 금액으로 간주
            fld = f"{grp}_amt" if grp in ("mat", "lab", "exp", "tot") else None
        elif grp and sub and not subk:
            fld = None
        else:
            label = top or sub
            if two_rows and sub and not top:
                label = sub
            for name, words in SINGLES:
                if _match(label, words):
                    fld = name
                    break
            # 2행 헤더에서 단가/금액 단독(그룹 없음) → 합계로 간주
            if two_rows and fld in ("price", "amt") and top and not grp:
                fld = f"tot_{fld}"
        if fld and fld in taken:
            # 중복이면 뒤의 것은 무시 (단, 같은 그룹 내 두 번째 '계' 등)
            cm.unmapped_cols.append(c)
            continue
        if fld:
            cm.fields[fld] = c
            taken.add(fld)
        else:
            cm.unmapped_cols.append(c)

    # 단일 단가/금액 양식에서 tot 로 승격
    if "amt" in cm.fields and "tot_amt" not in cm.fields and not cm.has_cost_split():
        cm.fields["tot_amt"] = cm.fields.pop("amt")
    if "price" in cm.fields and "tot_price" not in cm.fields and not cm.has_cost_split():
        cm.fields["tot_price"] = cm.fields.pop("price")

    essentials = ["name", "unit", "qty"]
    hit = sum(1 for e in essentials if e in cm.fields)
    amt_ok = any(k in cm.fields for k in ("tot_amt", "mat_amt", "lab_amt", "chg_amt"))
    cm.confidence = (hit / 3) * 0.7 + (0.3 if amt_ok else 0.0)
    return cm


def apply_user_mapping(cm: ColumnMap, mapping: dict[str, int | None]) -> ColumnMap:
    """사용자 지정 매핑 {field: col} 적용. None 은 해제."""
    for f, c in mapping.items():
        if c is None:
            cm.fields.pop(f, None)
        else:
            # 같은 열을 쓰던 다른 필드는 해제
            for k, v in list(cm.fields.items()):
                if v == c and k != f:
                    cm.fields.pop(k)
            cm.fields[f] = int(c)
    cm.unmapped_cols = [c for c in cm.raw_headers if c not in cm.fields.values()]
    return cm
