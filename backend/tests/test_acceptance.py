"""요청서 17장 합격 기준 테스트 1~8."""
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest

from tests.conftest import TEST_DIR, wait_done


def _upload(client, auth, pid, *paths):
    files = [("files", (p.name, p.read_bytes(), "application/octet-stream")) for p in paths]
    r = client.post(f"/api/projects/{pid}/files", headers=auth, files=files)
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture(scope="module")
def project(client, auth, fixtures):
    r = client.post("/api/projects", headers=auth, json={"name": "[테스트] 공원 CCTV 설치공사", "client": "테스트시청", "author": "테스터"})
    pid = r.json()["id"]
    _upload(client, auth, pid, fixtures / "test_cctv_plan.dxf", fixtures / "test_cctv_riser.dxf")
    files = wait_done(client, auth, pid)
    assert all(f["status"] == "done" for f in files), files
    return pid


def test_1_legend_recognition(client, auth, project):
    """서로 다른 CCTV 기호 2종 → 서로 다른 품목으로 식별."""
    r = client.post(f"/api/projects/{project}/legends/auto?wait=1", headers=auth)
    assert r.status_code == 200, r.text
    tables = client.get(f"/api/projects/{project}/legends", headers=auth).json()
    assert tables, "범례표가 자동 검출되지 않음"
    syms = [s for t in tables for s in t["symbols"]]
    cams = {s["block_name"]: s for s in syms if s["block_name"] in ("CCTV_FIX", "CCTV_PTZ")}
    assert set(cams) == {"CCTV_FIX", "CCTV_PTZ"}
    assert cams["CCTV_FIX"]["item_name"] != cams["CCTV_PTZ"]["item_name"]
    assert "고정형" in cams["CCTV_FIX"]["item_name"] and "PTZ" in cams["CCTV_PTZ"]["item_name"]
    assert cams["CCTV_FIX"]["fingerprint"]["hash"] != cams["CCTV_PTZ"]["fingerprint"]["hash"]
    assert cams["CCTV_FIX"]["unit"] == "대"
    # 선형 기호: 전선관/케이블 범주 + 레이어 힌트
    line = {s["item_name"]: s for s in syms if s["category"] in ("conduit", "cable")}
    assert "통신용 전선관" in line and line["통신용 전선관"]["layer_hint"] == "T-CONDUIT"
    assert line["LAN 케이블"]["category"] == "cable"
    client.post(f"/api/projects/{project}/legends/approve-all", headers=auth)


def test_2_quantity_detection(client, auth, project):
    """고정형 12대, PTZ 3대 — 범례표 안 참조 기호·기존 시설 제외."""
    r = client.post(f"/api/projects/{project}/takeoff/run?wait=1", headers=auth)
    assert r.status_code == 200, r.text
    items = client.get(f"/api/projects/{project}/takeoff/items", headers=auth).json()
    by = {(q["item_name"], q["install_type"], q["drawing_nos"][0] if q["drawing_nos"] else None): q for q in items}
    fix_new = [q for q in items if "고정형" in q["item_name"] and q["install_type"] == "new" and "C-01" in q["drawing_nos"]]
    ptz_new = [q for q in items if "PTZ" in q["item_name"] and q["install_type"] == "new" and "C-01" in q["drawing_nos"]]
    assert fix_new and fix_new[0]["qty_total"] == 12, fix_new
    assert ptz_new and ptz_new[0]["qty_total"] == 3, ptz_new
    existing = [q for q in items if "고정형" in q["item_name"] and q["install_type"] == "existing"]
    assert existing and existing[0]["qty_total"] == 1  # 기존 시설은 별도 구분
    # 근거: 검출 ID → 도면 좌표
    assert all(e["detected_symbol_id"] for e in fix_new[0]["evidence"])
    det = client.get(f"/api/projects/{project}/takeoff/symbols", headers=auth, params={"legend_symbol_id": fix_new[0]["legend_symbol_id"]}).json()
    assert len([d for d in det if d["install_type"] == "new" and d["sheet_id"] == fix_new[0]["evidence"][0]["sheet_id"]]) == 12
    assert all(d["x"] is not None and d["handle"] for d in det)


def test_2b_cross_validation(client, auth, project):
    """배치도 12대 vs 계통도 14대 → 임의 일치시키지 않고 검토 메시지."""
    issues = client.get(f"/api/projects/{project}/takeoff/issues", headers=auth).json()
    cross = [i for i in issues if i["code"] == "CROSS_CHECK"]
    assert cross, issues
    assert "차이" in cross[0]["message"] and "확인" in cross[0]["message"]


def test_3_conduit_length(client, auth, project):
    """mm 단위 10m+20m+15m 연결 구간 = 45m."""
    items = client.get(f"/api/projects/{project}/takeoff/items", headers=auth).json()
    cond = [q for q in items if q["item_name"] == "통신용 전선관"]
    assert cond, items
    assert abs(cond[0]["qty_drawing"] - 45.0) < 1e-6, cond[0]
    assert cond[0]["qty_total"] == 45.0
    routes = client.get(f"/api/projects/{project}/takeoff/routes", headers=auth).json()
    cr = [r for r in routes if r["kind"] == "conduit"]
    assert len(cr) == 1 and cr[0]["segment_count"] == 3 and abs(cr[0]["horizontal_m"] - 45) < 1e-6


def test_3b_duplicate_segment_detected(client, auth, fixtures):
    r = client.post("/api/projects", headers=auth, json={"name": "[테스트] 중복선분"})
    pid = r.json()["id"]
    _upload(client, auth, pid, fixtures / "test_cctv_plan_dup.dxf")
    wait_done(client, auth, pid)
    client.post(f"/api/projects/{pid}/legends/auto?wait=1", headers=auth)
    client.post(f"/api/projects/{pid}/legends/approve-all", headers=auth)
    client.post(f"/api/projects/{pid}/takeoff/run?wait=1", headers=auth)
    issues = client.get(f"/api/projects/{pid}/takeoff/issues", headers=auth).json()
    assert any(i["code"] == "DUPLICATE_SEGMENT" for i in issues), issues
    routes = [r for r in client.get(f"/api/projects/{pid}/takeoff/routes", headers=auth).json() if r["kind"] == "conduit"]
    assert routes[0]["status"] == "review"
    assert abs(routes[0]["horizontal_m"] - 45) < 1e-6          # 중복분 기본 제외
    dup = [s for s in routes[0]["segments"] if s["dup_of"]][0]
    # 사용자가 "별도 선로" 로 확인하면 포함
    r = client.put(f"/api/projects/{pid}/takeoff/segments/{dup['id']}", headers=auth, params={"status": "verified"})
    assert abs(r.json()["horizontal_m"] - 55) < 1e-6


def test_4_cable_circuits(client, auth, project):
    """경로 30m(여기선 45m) × 3회선 = 케이블 135m, 전선관은 45m 유지."""
    routes = client.get(f"/api/projects/{project}/takeoff/routes", headers=auth).json()
    route = [r for r in routes if r["kind"] == "conduit"][0]
    syms = [s for t in client.get(f"/api/projects/{project}/legends", headers=auth).json() for s in t["symbols"]]
    cable = [s for s in syms if s["category"] == "cable"][0]
    r = client.post(f"/api/projects/{project}/takeoff/circuits", headers=auth, json={"route_id": route["id"], "cable_symbol_id": cable["id"], "circuit_count": 3})
    assert r.status_code == 200, r.text
    items = client.get(f"/api/projects/{project}/takeoff/items", headers=auth).json()
    cab = [q for q in items if q["item_name"] == "LAN 케이블"][0]
    cond = [q for q in items if q["item_name"] == "통신용 전선관"][0]
    assert abs(cab["qty_total"] - 135.0) < 1e-6, cab
    assert abs(cond["qty_total"] - 45.0) < 1e-6
    assert "× 3회선" in cab["formula"]


def test_5_edit_and_recompute(client, auth, project):
    """잘못 인식된 CCTV 1대 제외 → 집계 즉시 갱신, 재분석 후에도 보존."""
    det = client.get(f"/api/projects/{project}/takeoff/symbols", headers=auth).json()
    fix = [d for d in det if d["install_type"] == "new" and d["method"] == "block"]
    syms = {s["id"]: s for t in client.get(f"/api/projects/{project}/legends", headers=auth).json() for s in t["symbols"]}
    target = [d for d in fix if syms[d["legend_symbol_id"]]["block_name"] == "CCTV_FIX"][0]
    r = client.put(f"/api/projects/{project}/takeoff/symbols/{target['id']}", headers=auth, json={"status": "excluded", "note": "오인식"})
    assert r.status_code == 200
    items = client.get(f"/api/projects/{project}/takeoff/items", headers=auth).json()
    fix_new = [q for q in items if "고정형" in q["item_name"] and q["install_type"] == "new" and "C-01" in q["drawing_nos"]][0]
    assert fix_new["qty_total"] == 11
    # 자동 재분석이 사용자 수정을 초기화하지 않음
    client.post(f"/api/projects/{project}/takeoff/run?wait=1", headers=auth)
    items = client.get(f"/api/projects/{project}/takeoff/items", headers=auth).json()
    fix_new = [q for q in items if "고정형" in q["item_name"] and q["install_type"] == "new" and "C-01" in q["drawing_nos"]][0]
    assert fix_new["qty_total"] == 11
    # 누락 기호 수동 추가 → 12
    r = client.post(f"/api/projects/{project}/takeoff/symbols/manual", headers=auth,
                    json={"sheet_id": target["sheet_id"], "legend_symbol_id": target["legend_symbol_id"], "x": 100, "y": 100})
    assert r.status_code == 200
    items = client.get(f"/api/projects/{project}/takeoff/items", headers=auth).json()
    fix_new = [q for q in items if "고정형" in q["item_name"] and q["install_type"] == "new" and "C-01" in q["drawing_nos"]][0]
    assert fix_new["qty_total"] == 12


def test_6_excel_export(client, auth, project):
    r = client.get(f"/api/projects/{project}/export/xlsx", headers=auth)
    assert r.status_code == 200
    out = TEST_DIR / "export.xlsx"
    out.write_bytes(r.content)
    import openpyxl
    wb = openpyxl.load_workbook(out)
    names = wb.sheetnames
    assert len(names) == 10 and names[0].startswith("1.") and "4.상세산출서" in names and "10.범례매칭표" in names
    ws = wb["4.상세산출서"]
    vals = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert any(v == "통신용 전선관" for v in vals) and any(v == "산출식" for v in vals)
    assert ws.print_title_rows and ws.page_setup.orientation == "landscape"
    r = client.get(f"/api/projects/{project}/export/pdf", headers=auth)
    assert r.status_code == 200 and r.content[:4] == b"%PDF"


def test_7_persistence_after_restart(client, auth, project):
    """실제 uvicorn 서버 기동 → 재시작 → 동일 계정 로그인 → 도면·산출 결과 유지."""
    import httpx
    env = {**os.environ, "KET_DATABASE_URL": os.environ["KET_DATABASE_URL"], "KET_STORAGE_DIR": os.environ["KET_STORAGE_DIR"]}
    cwd = str(Path(__file__).resolve().parents[1])

    def start():
        p = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.main:app", "--port", "8765"], cwd=cwd, env=env,
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        for _ in range(60):
            try:
                if httpx.get("http://127.0.0.1:8765/api/health").status_code == 200:
                    return p
            except Exception:
                time.sleep(0.5)
        p.kill()
        raise RuntimeError("server did not start")

    def login_and_fetch():
        tok = httpx.post("http://127.0.0.1:8765/api/auth/login", data={"username": "tester@example.com", "password": "password123"}).json()["access_token"]
        h = {"Authorization": f"Bearer {tok}"}
        files = httpx.get(f"http://127.0.0.1:8765/api/projects/{project}/files", headers=h).json()
        items = httpx.get(f"http://127.0.0.1:8765/api/projects/{project}/takeoff/items", headers=h).json()
        return files, items

    p = start()
    try:
        f1, i1 = login_and_fetch()
    finally:
        p.terminate()
        p.wait(10)
    p = start()
    try:
        f2, i2 = login_and_fetch()
    finally:
        p.terminate()
        p.wait(10)
    assert len(f1) == len(f2) == 2 and all(f["status"] == "done" and f["sheets"] for f in f2)
    assert [(q["item_name"], q["qty_total"]) for q in i1] == [(q["item_name"], q["qty_total"]) for q in i2]
    assert any(q["item_name"] == "통신용 전선관" and q["qty_total"] == 45 for q in i2)


def test_8_scanned_pdf_requires_scale(client, auth, fixtures):
    """축척 미상 스캔 PDF → 길이 산출 대신 축척 확인 요청."""
    r = client.post("/api/projects", headers=auth, json={"name": "[테스트] 스캔PDF"})
    pid = r.json()["id"]
    _upload(client, auth, pid, fixtures / "test_scanned_noscale.pdf")
    files = wait_done(client, auth, pid)
    f = files[0]
    assert f["status"] == "done", f
    sheet = f["sheets"][0]
    assert sheet["source_kind"] == "pdf_raster" and sheet["scale_confirmed"] is False and sheet["scale_factor"] is None
    assert any("축척" in w for w in f["warnings"])
    # 사용자가 레이어 힌트 없는 선형 범례를 등록하고 산출 실행해도 길이 산출 불가 + 이슈
    tables = client.post(f"/api/projects/{pid}/legends/region", headers=auth,
                         json={"sheet_id": sheet["id"], "region": {"x0": 0, "y0": 0, "x1": 100, "y1": 100}, "title": "수동"}).json()
    client.post(f"/api/projects/{pid}/legends/{tables['id']}/symbols", headers=auth,
                json={"symbol_name": "CD", "item_name": "전선관", "spec": "CD 22", "unit": "m", "category": "conduit", "layer_hint": "PDF"})
    client.post(f"/api/projects/{pid}/takeoff/run?wait=1", headers=auth)
    issues = client.get(f"/api/projects/{pid}/takeoff/issues", headers=auth).json()
    assert any(i["code"] == "SCALE_UNCONFIRMED" for i in issues), issues
    items = client.get(f"/api/projects/{pid}/takeoff/items", headers=auth).json()
    assert all(q["qty_total"] is None for q in items if q["unit"] == "m"), items
    # 축척 확정(기준거리 100pt = 10m) 후에는 길이 산출
    r = client.post(f"/api/projects/{pid}/sheets/{sheet['id']}/scale", headers=auth, json={"ref_points": [[0, 0], [100, 0]], "ref_length_m": 10})
    assert r.json()["scale_confirmed"] is True and abs(r.json()["scale_factor"] - 0.1) < 1e-9
    client.post(f"/api/projects/{pid}/takeoff/run?wait=1", headers=auth)
    items = client.get(f"/api/projects/{pid}/takeoff/items", headers=auth).json()
    assert all(q["qty_total"] is not None for q in items if q["unit"] == "m" and q["item_name"] == "전선관"), items


def test_9_vector_pdf_fingerprint(client, auth, fixtures):
    """벡터 PDF: LEGEND 표의 원형 기호 → 도면의 동일 도형 5개 검출 (범례 안 기호 제외)."""
    r = client.post("/api/projects", headers=auth, json={"name": "[테스트] 벡터PDF"})
    pid = r.json()["id"]
    _upload(client, auth, pid, fixtures / "test_vector.pdf")
    files = wait_done(client, auth, pid)
    assert files[0]["sheets"][0]["source_kind"] == "pdf_vector"
    client.post(f"/api/projects/{pid}/legends/auto?wait=1", headers=auth)
    tables = client.get(f"/api/projects/{pid}/legends", headers=auth).json()
    assert tables and tables[0]["symbols"], tables
    s = tables[0]["symbols"][0]
    assert s["item_name"] == "WALL LIGHT" and s["fingerprint"]["circles"] >= 1
    client.post(f"/api/projects/{pid}/legends/approve-all", headers=auth)
    client.post(f"/api/projects/{pid}/takeoff/run?wait=1", headers=auth)
    items = client.get(f"/api/projects/{pid}/takeoff/items", headers=auth).json()
    wl = [q for q in items if q["item_name"] == "WALL LIGHT"]
    assert wl and wl[0]["qty_total"] == 5, items


def test_10_security_company_isolation(client, auth, project):
    r = client.post("/api/auth/register", json={"company_name": "다른회사", "email": "other@example.com", "name": "타사", "password": "password123"})
    h = {"Authorization": f"Bearer {r.json()['access_token']}"}
    assert client.get(f"/api/projects/{project}", headers=h).status_code == 403
    assert client.get(f"/api/projects/{project}/takeoff/items").status_code == 401
    bad = client.post(f"/api/projects/{project}/files", headers=auth, files=[("files", ("evil.exe", b"MZ....", "application/octet-stream"))]).json()
    assert "error" in bad[0]
    fake = client.post(f"/api/projects/{project}/files", headers=auth, files=[("files", ("x.pdf", b"not a pdf", "application/pdf"))]).json()
    assert "error" in fake[0]


def test_11_estimate_stage5(client, auth, project):
    """5단계: 기준정보 등록 → 원가계산 → 미산정 표시 → 검증 전 확정 불가 → 검증 후 확정 → Excel."""
    import openpyxl
    client.post("/api/labor-rates", headers=auth, json={"trade": "통신내선공", "daily_rate": 300000, "publisher": "테스트협회", "effective_date": "2026-01-01"})
    client.post("/api/norms", headers=auth, json={"code": "T-001", "title": "CCTV 카메라 설치", "unit": "대", "labor": {"통신내선공": 0.5}, "publisher": "테스트원", "year": 2026})
    client.post("/api/norms", headers=auth, json={"code": "T-002", "title": "전선관 포설", "unit": "m", "labor": {"통신내선공": 0.02, "보통인부": 0.01}, "publisher": "테스트원", "year": 2026})
    m = client.post("/api/materials", headers=auth, json={"item_name": "고정형 네트워크 카메라", "spec": "4MP, 옥외형", "unit": "대", "supply_type": "gov", "price": 500000, "price_source": "업체견적", "effective_date": "2026-01-01"}).json()
    items = client.get(f"/api/projects/{project}/takeoff/items", headers=auth).json()
    cam = [q for q in items if "고정형" in q["item_name"] and q["install_type"] == "new"][0]
    cond = [q for q in items if q["item_name"] == "통신용 전선관"][0]
    body = {"rates": {"간접노무비": {"pct": 10, "base": "직접노무비", "source": "테스트 제비율"}, "이윤": {"pct": 15, "base": "직접공사비"}},
            "vat_pct": 10, "mapping": {str(cam["id"]): {"norm_code": "T-001", "material_id": m["id"]}, str(cond["id"]): {"norm_code": "T-002"}}}
    r = client.post(f"/api/projects/{project}/estimate", headers=auth, json=body)
    assert r.status_code == 200, r.text
    e = r.json()
    L = {x["item_id"]: x for x in e["lines"]}
    assert L[cam["id"]]["labor_cost"] == round(0.5 * 300000 * 12)
    assert L[cam["id"]]["material_cost"] == 500000 * 12 and "관급자재" in L[cam["id"]]["notes"]
    assert e["gov_material"] == 6000000 and e["direct_material"] == 0
    assert L[cond["id"]]["labor_cost"] is None and any("보통인부" in n for n in L[cond["id"]]["notes"])  # 노임 미등록 → 미산정
    assert e["overhead"]["이윤"]["amount"] is None  # 근거 없음 → 미산정
    assert e["overhead"]["간접노무비"]["amount"] == round(e["direct_labor"] * 0.1)
    assert e["finalizable"] is False
    assert client.post(f"/api/projects/{project}/estimate/{e['id']}/finalize", headers=auth).status_code == 400
    assert client.get(f"/api/projects/{project}/estimate", headers=auth).json()["id"] == e["id"]
    r = client.get(f"/api/projects/{project}/estimate/{e['id']}/xlsx", headers=auth)
    assert r.status_code == 200
    out = TEST_DIR / "estimate.xlsx"
    out.write_bytes(r.content)
    assert openpyxl.load_workbook(out).sheetnames[:2] == ["1.원가계산서", "2.일위대가"]
    # 전체 검증 후 확정 가능
    for q in items:
        client.put(f"/api/projects/{project}/takeoff/items/{q['id']}", headers=auth, json={"status": "verified"})
    e2 = client.post(f"/api/projects/{project}/estimate", headers=auth, json=body).json()
    assert e2["finalizable"] is True
    assert client.post(f"/api/projects/{project}/estimate/{e2['id']}/finalize", headers=auth).json()["ok"]


def test_12_refdata_excel_import(client, auth):
    import io
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["품명", "규격", "단위", "관급/사급", "단가", "출처", "적용일"]); ws.append(["LAN 케이블", "CAT.6 4P", "m", "사급", 1200, "물가정보 2026-01", "2026-01-01"])
    b = io.BytesIO(); wb.save(b)
    r = client.post("/api/refdata/materials/import", headers=auth, files=[("file", ("m.xlsx", b.getvalue(), "application/octet-stream"))])
    assert r.status_code == 200 and r.json()["imported"] == 1
    assert any(x["item_name"] == "LAN 케이블" and x["price"] == 1200 for x in client.get("/api/materials", headers=auth).json())


def test_13_background_job(client, auth, project):
    """웹 타임아웃 방지: 산출은 백그라운드 작업으로 실행되고 진행률을 조회할 수 있어야 한다."""
    r = client.post(f"/api/projects/{project}/takeoff/run", headers=auth).json()
    assert "job_id" in r
    for _ in range(120):
        j = client.get(f"/api/projects/{project}/jobs/{r['job_id']}", headers=auth).json()
        if j["status"] in ("done", "failed"):
            break
        time.sleep(0.5)
    assert j["status"] == "done", j
    assert j["progress"] == 100 and j["result"]["items"] > 0


def test_14_dwg_upload_converts(client, auth):
    """DWG 업로드 → 내장 LibreDWG 로 DXF 변환 → 시트·레이어·블록 생성. 변환기가 없으면 '변환기 필요' 로 명확히 표시."""
    from app.engine.dxf_parser import find_dwg2dxf
    r = client.post("/api/projects", headers=auth, json={"name": "[테스트] DWG 업로드"})
    pid = r.json()["id"]
    samples = Path(__file__).parent / "dwg_samples"
    _upload(client, auth, pid, samples / "example_2018.dwg", samples / "example_2007.dwg")
    files = wait_done(client, auth, pid)
    for f in files:
        if find_dwg2dxf() is None:
            assert f["status"] == "needs_converter" and "변환기" in f["error"]
            continue
        assert f["status"] == "done", f
        s = f["sheets"][0]
        assert s["source_kind"] == "dxf" and s["units"] == "mm" and s["scale_confirmed"]
        assert s["stats"]["inserts"] == 10 and s["stats"]["entities"] >= 40
        assert any("LibreDWG" in w for w in f["warnings"])


def test_15_reprocess_after_takeoff(client, auth, project):
    """산출·범례 승인 후 '재분석' 해도 FK 오류 없이 다시 처리되고, 승인된 범례는 보존되며 재산출이 된다."""
    client.post(f"/api/projects/{project}/takeoff/run?wait=1", headers=auth)
    n_leg = len(client.get(f"/api/projects/{project}/legends", headers=auth).json())
    files = client.get(f"/api/projects/{project}/files", headers=auth).json()
    for f in files:
        r = client.post(f"/api/projects/{project}/files/{f['id']}/reprocess", headers=auth)
        assert r.status_code == 200, r.text
    files = wait_done(client, auth, project)
    assert all(f["status"] == "done" for f in files), [(f["name"], f["error"]) for f in files]
    assert len(client.get(f"/api/projects/{project}/legends", headers=auth).json()) == n_leg
    r = client.post(f"/api/projects/{project}/takeoff/run?wait=1", headers=auth).json()
    assert r["items"] > 0


def test_16_region_tally_manual(client, auth, fixtures):
    """자동 인식 없이: 드래그 영역 안 객체를 모양별로 세고, 케이블 길이는 직접 입력 → 물량표에 합산.
    DXF(블록) 와 벡터 PDF(닫힌 도형) 모두."""
    pid = client.post("/api/projects", headers=auth, json={"name": "[테스트] 영역 집계"}).json()["id"]
    _upload(client, auth, pid, fixtures / "test_cctv_plan.dxf", fixtures / "test_vector.pdf")
    files = {f["name"]: f for f in wait_done(client, auth, pid)}
    dxf = files["test_cctv_plan.dxf"]["sheets"][0]
    # 배치도 영역(범례표 제외): 고정형 12 + 기존 1, PTZ 3, 접속함 1
    r = client.post(f"/api/projects/{pid}/sheets/{dxf['id']}/region-analyze", headers=auth,
                    json={"region": {"x0": 0, "y0": 0, "x1": 30000, "y1": 30000}}).json()
    blocks = {o["block"]: o["count"] for o in r["objects"] if o["kind"] == "block"}
    assert blocks == {"CCTV_FIX": 13, "CCTV_PTZ": 3, "JB_OUT": 1}, blocks
    conduit = next(l for l in r["lines"] if l["key"] == "T-CONDUIT")
    assert abs(conduit["length"] - 45000) < 1          # 10+20+15 m (mm 단위)
    # 기존 시설 레이어를 숨기고 다시 → 고정형 12
    r2 = client.post(f"/api/projects/{pid}/sheets/{dxf['id']}/region-analyze", headers=auth,
                     json={"region": {"x0": 0, "y0": 0, "x1": 30000, "y1": 30000}, "hidden_layers": ["T-EXIST"]}).json()
    assert {o["block"]: o["count"] for o in r2["objects"] if o["kind"] == "block"}["CCTV_FIX"] == 12
    # 벡터 PDF: 원+가로선 기호 5개가 한 그룹
    pdf = files["test_vector.pdf"]["sheets"][0]
    H = pdf["bbox"]["y1"]
    r3 = client.post(f"/api/projects/{pid}/sheets/{pdf['id']}/region-analyze", headers=auth,
                     json={"region": {"x0": 100, "y0": H - 340, "x1": 600, "y1": H - 260}}).json()
    assert r3["mode"] == "vector" and r3["objects"] and r3["objects"][0]["count"] == 5, [(o["label"], o["count"]) for o in r3["objects"]]
    # 저장 → 물량표 (같은 품명·규격 합산, 케이블은 입력 길이)
    rows = [{"kind": "object", "item_name": "고정형 카메라", "spec": "4MP", "qty": 12},
            {"kind": "cable", "item_name": "UTP 케이블", "spec": "CAT.6", "qty": 52.5}]
    assert client.post(f"/api/projects/{pid}/tallies", headers=auth,
                       json={"sheet_id": dxf["id"], "region": {"x0": 0, "y0": 0, "x1": 30000, "y1": 30000}, "rows": rows}).status_code == 200
    client.post(f"/api/projects/{pid}/tallies", headers=auth,
                json={"sheet_id": pdf["id"], "region": {"x0": 100, "y0": 1, "x1": 600, "y1": 300}, "rows": [{"item_name": "고정형 카메라", "spec": "4MP", "qty": 5}]})
    items = {(q["item_name"], q["spec"]): q for q in client.get(f"/api/projects/{pid}/takeoff/items", headers=auth).json()}
    cam = items[("고정형 카메라", "4MP")]
    assert cam["qty_total"] == 17 and cam["unit"] == "개" and cam["status"] == "verified" and "영역#" in cam["formula"]
    assert items[("UTP 케이블", "CAT.6")]["qty_total"] == 52.5 and items[("UTP 케이블", "CAT.6")]["unit"] == "m"
    # 자동 산출을 다시 돌려도 수동 집계는 유지, 삭제하면 빠짐
    client.post(f"/api/projects/{pid}/takeoff/run?wait=1", headers=auth)
    tl = client.get(f"/api/projects/{pid}/tallies", headers=auth).json()
    assert len(tl) == 3
    client.delete(f"/api/projects/{pid}/tallies/{tl[-1]['id']}", headers=auth)
    items = {(q["item_name"], q["spec"]): q for q in client.get(f"/api/projects/{pid}/takeoff/items", headers=auth).json()}
    assert items[("고정형 카메라", "4MP")]["qty_total"] == 12
    assert client.get(f"/api/projects/{pid}/export/xlsx", headers=auth).status_code == 200


def test_17_symbol_library_suggest(client, auth, tmp_path):
    """학습한 표준 심볼 라이브러리: 블록명 일치 → 이름 추천, 라이브러리 목록 API."""
    import ezdxf
    lib = client.get("/api/symbol-library", headers=auth).json()
    assert len(lib["symbols"]) >= 200 and any("TTA" in s for s in lib["sources"])
    doc = ezdxf.new("R2018"); doc.header["$INSUNITS"] = 4
    b = doc.blocks.new("Z-죠인트박스"); b.add_circle((0, 0), 100); b.add_text("J", dxfattribs={"height": 80})
    msp = doc.modelspace()
    for i in range(6):
        msp.add_blockref("Z-죠인트박스", (1000 + i * 1000, 1000))
    msp.add_text("LIGHTING PLAN", dxfattribs={"height": 200}).set_placement((0, 3000))
    p = tmp_path / "jb.dxf"; doc.saveas(p)
    pid = client.post("/api/projects", headers=auth, json={"name": "[테스트] 심볼 추천"}).json()["id"]
    _upload(client, auth, pid, p)
    sh = wait_done(client, auth, pid)[0]["sheets"][0]
    r = client.post(f"/api/projects/{pid}/sheets/{sh['id']}/region-analyze", headers=auth,
                    json={"region": {"x0": 0, "y0": 0, "x1": 8000, "y1": 2000}}).json()
    o = r["objects"][0]
    assert o["count"] == 6 and o["suggest"][0]["name"] == "조인트박스(J/B)" and o["suggest"][0]["score"] > 0.9


def test_18_legend_delete_and_guard(client, auth, fixtures):
    """범례 삭제(검출·산출 후에도 FK 오류 없이), 도면 대부분을 덮는 범례 영역 차단, 도면 전체 집계는 범례 견본 제외."""
    pid = client.post("/api/projects", headers=auth, json={"name": "[테스트] 범례 삭제"}).json()["id"]
    _upload(client, auth, pid, fixtures / "test_cctv_plan.dxf")
    sh = wait_done(client, auth, pid)[0]["sheets"][0]
    client.post(f"/api/projects/{pid}/legends/auto?wait=1", headers=auth)
    client.post(f"/api/projects/{pid}/legends/approve-all", headers=auth)
    client.post(f"/api/projects/{pid}/takeoff/run?wait=1", headers=auth)
    assert client.get(f"/api/projects/{pid}/takeoff/items", headers=auth).json()
    b = sh["bbox"]
    # 전체 집계: 범례표 안 견본 블록(CCTV_FIX 등 3개)은 빠져야 함 → 고정형 12+기존 1
    r = client.post(f"/api/projects/{pid}/sheets/{sh['id']}/region-analyze", headers=auth, json={"region": b}).json()
    assert {o["block"]: o["count"] for o in r["objects"] if o["kind"] == "block"}["CCTV_FIX"] == 13
    # 도면 전체를 범례로 지정 → 400 (force 없이)
    big = client.post(f"/api/projects/{pid}/legends/region", headers=auth, json={"sheet_id": sh["id"], "region": b, "title": "x"})
    assert big.status_code == 400 and "범례" in big.json()["detail"]
    # 범례 한 행 삭제 → 전체 삭제 (검출·산출이 있어도 성공), 산출 수량 정리
    legs = client.get(f"/api/projects/{pid}/legends", headers=auth).json()
    sid = legs[0]["symbols"][0]["id"]
    assert client.delete(f"/api/projects/{pid}/legends/symbols/{sid}", headers=auth).status_code == 200
    assert client.delete(f"/api/projects/{pid}/legends", headers=auth).json()["deleted"] == len(legs)
    assert client.get(f"/api/projects/{pid}/legends", headers=auth).json() == []
    assert client.get(f"/api/projects/{pid}/takeoff/items", headers=auth).json() == []


def test_19_damaged_dxf_recovers(client, auth, tmp_path):
    """EOF 누락 + 이름 없는 레이어 같은 손상 DXF 도 보정해서 읽고, 경고로 알린다."""
    import ezdxf
    doc = ezdxf.new(); doc.layers.add("E-LIGHT"); msp = doc.modelspace()
    for i in range(5):
        msp.add_circle((i * 100, 0), 20, dxfattribs={"layer": "E-LIGHT"})
    p = tmp_path / "bad.dxf"; doc.saveas(p)
    t = p.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
    for k in range(len(t) - 1):
        if t[k].strip() == "2" and t[k + 1] == "E-LIGHT":
            del t[k:k + 2]
            break
    s = "\n".join(t); p.write_text(s[:s.rfind("EOF") - 4], encoding="utf-8")
    pid = client.post("/api/projects", headers=auth, json={"name": "[테스트] 손상 DXF"}).json()["id"]
    _upload(client, auth, pid, p)
    f = wait_done(client, auth, pid)[0]
    assert f["status"] == "done", f["error"]
    assert any("보정" in w for w in f["warnings"])
    assert f["sheets"][0]["stats"]["entities"] >= 5
