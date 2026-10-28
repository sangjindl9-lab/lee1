"""관제점(태그) 맵.

현장마다 PLC/RTU 주소가 다르므로 프로그램은 '태그 이름'으로만 동작하고,
태그 이름 → Modbus 주소 연결은 이 맵(기본값) 또는 현장 관제점 일람표 엑셀(tags/<터널키>.json)로 정한다.

구분(type)
  DI : Discrete Input  (FC2, 1xxxx)  - 상태 입력
  DO : Coil            (FC1/5, 0xxxx) - 제어 출력
  AI : Input Register  (FC4, 3xxxx)  - 계측 입력
  AO : Holding Register(FC3/6, 4xxxx) - 설정값
  PI : Holding Register 2워드(uint32) - 적산값
주소(addr)는 0부터 시작하는 오프셋(프로토콜 주소)이다. 일람표의 30001 같은 표기는 불러올 때 자동 변환한다.
"""
import csv
import io
import json
import os

STAGES = [
    {"k": "맑음", "ent": 100, "base": 100},
    {"k": "흐림", "ent": 50, "base": 100},
    {"k": "우천", "ent": 25, "base": 100},
    {"k": "야간", "ent": 0, "base": 100},
    {"k": "심야", "ent": 0, "base": 50},
]

FEEDER_USE = {
    "LV-101": "기본조명 (LTR-1)", "LV-102": "입구부 조명", "LV-103": "제트팬 MCC",
    "LV-104": "배수펌프", "LV-105": "방재설비", "LV-106": "소화설비", "LV-107": "전열·환기",
    "LV-108": "관리동", "LV-109": "예비", "LV-UPS": "UPS",
}

FC = {"DO": 1, "DI": 2, "AO": 3, "PI": 3, "AI": 4}
TYPES = set(FC)
MAX_FANS = 10
MAX_FEEDERS = 10


def fan_ids(n):
    return ["JF-%02d" % (i + 1) for i in range(n)]


def default_tags(t):
    """터널 설정(t)으로 기본 관제점 맵을 만든다. 시뮬레이터는 항상 이 맵을 쓴다."""
    T = []

    def add(tag, desc, typ, addr, scale=1.0, unit="", dtype=None):
        if dtype is None:
            dtype = "bool" if typ in ("DI", "DO") else ("uint32" if typ == "PI" else "int16")
        T.append({"tag": tag, "desc": desc, "type": typ, "addr": addr,
                  "scale": scale, "unit": unit, "dtype": dtype})

    feeders = t["feeders"][:MAX_FEEDERS]
    fans = fan_ids(min(t["fans"], MAX_FANS))

    # ---------- DI ----------
    add("HV-1.VCB", "수전 VCB 투입 상태", "DI", 0)
    add("HV-1.OCR", "과전류계전기 동작", "DI", 1)
    add("HV-1.OCGR", "지락과전류계전기 동작", "DI", 2)
    add("HV-1.UVR", "부족전압계전기 동작 (정전)", "DI", 3)
    add("PTR-1.TALM", "PTR-1 온도 경보", "DI", 4)
    for i, f in enumerate(feeders):
        add(f + ".ST", FEEDER_USE.get(f, f) + " 차단기 투입 상태", "DI", 10 + i * 2)
        add(f + ".TRIP", FEEDER_USE.get(f, f) + " 트립", "DI", 11 + i * 2)
    add("UPS.RUN", "UPS 운전", "DI", 30)
    add("UPS.BATT_DIS", "UPS 축전지 방전 중", "DI", 31)
    for i, s in enumerate(STAGES):
        add("LT.ST%d_FB" % (i + 1), "조명 %s 단계 응답" % s["k"], "DI", 40 + i)
    for i, f in enumerate(fans):
        b = 50 + i * 5
        add(f + ".FWD", f + " 정회전 운전", "DI", b)
        add(f + ".REV", f + " 역회전 운전", "DI", b + 1)
        add(f + ".FLT", f + " 고장 (과부하)", "DI", b + 2)
        add(f + ".REM", f + " 원격 선택", "DI", b + 3)
        add(f + ".WAIT", f + " 방향전환 대기", "DI", b + 4)
    for z in range(4):
        add("FD.Z%d" % (z + 1), "화재감지 구역 %d" % (z + 1), "DI", 110 + z)
    add("EV.PWR", "유도등 전원", "DI", 115)

    # ---------- DO ----------
    add("HV-1.CLOSE", "수전 VCB 투입 명령 (펄스)", "DO", 0)
    add("HV-1.OPEN", "수전 VCB 개방 명령 (펄스)", "DO", 1)
    for i, f in enumerate(feeders):
        add(f + ".CLOSE", f + " 투입 명령 (펄스)", "DO", 10 + i * 2)
        add(f + ".OPEN", f + " 개방 명령 (펄스)", "DO", 11 + i * 2)
    for i, s in enumerate(STAGES):
        add("LT.ST%d_CMD" % (i + 1), "조명 %s 단계 명령" % s["k"], "DO", 40 + i)
    for i, f in enumerate(fans):
        b = 50 + i * 5
        add(f + ".FWD_CMD", f + " 정회전 명령", "DO", b)
        add(f + ".REV_CMD", f + " 역회전 명령", "DO", b + 1)
        add(f + ".RESET", f + " 고장 복귀 (펄스)", "DO", b + 2)
    add("EV.BLINK", "유도등 점멸 명령", "DO", 115)
    add("EV.GUIDE", "유도등 피난유도 명령", "DO", 116)
    add("FD.RESET", "화재 수신기 복구 (펄스)", "DO", 117)

    # ---------- AI ----------
    for k, p in enumerate(["RS", "ST", "TR"]):
        add("HV-1.V_" + p, "수전 선간전압 " + p, "AI", k, 0.01, "kV")
    for k, p in enumerate(["R", "S", "T"]):
        add("HV-1.A_" + p, "수전 전류 %s상" % p, "AI", 3 + k, 0.1, "A")
    add("HV-1.KW", "수전 유효전력", "AI", 6, 0.1, "kW")
    add("HV-1.PF", "수전 역률", "AI", 7, 0.001, "")
    add("HV-1.HZ", "주파수", "AI", 8, 0.1, "Hz")
    add("PTR-1.TEMP", "PTR-1 권선온도", "AI", 9, 0.1, "℃")
    add("LTR-1.TEMP", "LTR-1 권선온도", "AI", 10, 0.1, "℃")
    for i, f in enumerate(feeders):
        add(f + ".A", f + " 전류", "AI", 20 + i * 2, 0.1, "A")
        add(f + ".KW", f + " 유효전력", "AI", 21 + i * 2, 0.1, "kW")
    add("UPS.BATT", "UPS 축전지 잔량", "AI", 40, 1, "%")
    add("LT.L20", "외부휘도 L20", "AI", 41, 1, "cd/m²")
    add("ENV.CO1", "CO 농도 #1", "AI", 42, 0.1, "ppm")
    add("ENV.CO2", "CO 농도 #2", "AI", 43, 0.1, "ppm")
    add("ENV.VI1", "VI 소광계수 #1", "AI", 44, 0.00001, "m⁻¹")
    add("ENV.VI2", "VI 소광계수 #2", "AI", 45, 0.00001, "m⁻¹")
    add("ENV.WS", "풍속", "AI", 46, 0.1, "m/s")
    add("ENV.WD", "풍향 (1=종점방향)", "AI", 47, 1, "")
    add("ENV.TRAFFIC", "교통량 검지", "AI", 48, 1, "대/h")
    add("FD.POS", "화재 위치 (선형감지기, ‰)", "AI", 49, 1, "‰")
    for i, f in enumerate(fans):
        add(f + ".A", f + " 운전 전류", "AI", 50 + i, 0.1, "A")

    # ---------- AO / PI ----------
    add("HV-1.KWH", "수전 적산전력량", "PI", 0, 1, "kWh")
    add("EV.SPLIT", "피난유도 방향 분기점 (‰)", "AO", 10, 1, "‰")
    for i, f in enumerate(fans):
        add(f + ".HRS", f + " 누적 가동시간", "PI", 20 + i * 2, 0.1, "h")

    # ---------- 시뮬레이터 전용 (실제 PLC에는 없음) ----------
    add("SIM.FIRE", "[모의] 화재 위치 ‰ (0=없음)", "AO", 100)
    add("SIM.FAULT", "[모의] 제트팬 고장 번호", "AO", 101)
    add("SIM.KEPCO", "[모의] 한전 수전 (1=정상)", "AO", 102)
    add("SIM.CO_BOOST", "[모의] CO 상승량", "AO", 103)
    add("SIM.SPEED", "[모의] 시간 배속", "AO", 104)
    add("SIM.MINUTE", "[모의] 현장시각 (분)", "AO", 105)
    add("SIM.DAY", "[모의] 경과 일수", "AO", 106)
    return T


# ---------- 일람표 불러오기 / 내보내기 ----------
HEAD = ["태그", "설명", "구분", "주소", "배율", "단위", "형식"]


def _norm_addr(typ, raw):
    """일람표 주소 표기(10001, 30001, 40001, 00001 등)를 0부터 시작하는 오프셋으로 바꾼다."""
    s = str(raw).strip()
    if not s:
        raise ValueError("주소가 비어 있습니다")
    v = int(float(s))
    if len(s) >= 5:
        lead = {"DO": 0, "DI": 1, "AI": 3, "AO": 4, "PI": 4}[typ]
        if v // 10000 == lead and v % 10000 >= 1:
            return v % 10000 - 1
    return v


def parse_rows(rows):
    """[{머리글: 값}] → 태그 목록. 머리글은 한글/영문 혼용 허용."""
    def pick(r, keys):
        for k, v in r.items():
            kk = str(k).replace(" ", "").upper()
            if any(x in kk for x in keys):
                return "" if v is None else str(v).strip()
        return ""

    out, errs = [], []
    for n, r in enumerate(rows, start=2):
        tag = pick(r, ["태그", "TAG"])
        if not tag:
            continue
        typ = (pick(r, ["구분", "유형", "TYPE"]) or "DI").upper()
        if typ not in TYPES:
            errs.append("%d행 %s: 구분 '%s'는 DI/DO/AI/AO/PI 중 하나여야 합니다" % (n, tag, typ))
            continue
        try:
            addr = _norm_addr(typ, pick(r, ["주소", "ADDR"]))
        except ValueError as e:
            errs.append("%d행 %s: %s" % (n, tag, e))
            continue
        sc = pick(r, ["배율", "SCALE"])
        dt = pick(r, ["형식", "DTYPE"]).lower()
        out.append({
            "tag": tag, "desc": pick(r, ["설명", "명칭", "DESC"]), "type": typ, "addr": addr,
            "scale": float(sc) if sc else 1.0, "unit": pick(r, ["단위", "UNIT"]),
            "dtype": dt or ("bool" if typ in ("DI", "DO") else "uint32" if typ == "PI" else "int16"),
        })
    return out, errs


def read_table_file(name, data):
    if name.lower().endswith(".csv"):
        txt = data.decode("utf-8-sig", errors="replace")
        return list(csv.DictReader(io.StringIO(txt)))
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [str(h or "").strip() for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]


def display_addr(t):
    lead = {"DO": 0, "DI": 1, "AI": 3, "AO": 4, "PI": 4}[t["type"]]
    return "%d%04d" % (lead, t["addr"] + 1)


def to_xlsx(tags, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "관제점일람표"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3EA")
        c.alignment = Alignment(horizontal="center")
    for t in tags:
        ws.append([t["tag"], t["desc"], t["type"], display_addr(t), t["scale"], t["unit"], t["dtype"]])
    for col, w in zip("ABCDEFG", [18, 34, 7, 10, 9, 8, 8]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def map_path(base, key):
    return os.path.join(base, "tags", key + ".json")


def load_tags(base, t):
    p = map_path(base, t["key"])
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            custom = json.load(f)
        # 시뮬레이터 전용 태그는 현장 맵에 없어도 유지
        names = {x["tag"] for x in custom}
        return custom + [x for x in default_tags(t) if x["tag"].startswith("SIM.") and x["tag"] not in names], True
    return default_tags(t), False


def save_tags(base, key, tags):
    os.makedirs(os.path.join(base, "tags"), exist_ok=True)
    with open(map_path(base, key), "w", encoding="utf-8") as f:
        json.dump(tags, f, ensure_ascii=False, indent=1)


def reset_tags(base, key):
    p = map_path(base, key)
    if os.path.exists(p):
        os.remove(p)
