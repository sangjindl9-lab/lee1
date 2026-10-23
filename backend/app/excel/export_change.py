"""설계변경 엑셀 출력: 증감총괄 / 변경 원가계산서 / 설계변경 대비표(좌우 3블록) / 변경내역서(당초·변경 2행형) / 신규품목 내역서."""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from ..services.design_change import ChangeLine
from .export_progress import BOLD, FONT, HEAD_FILL, NUM, NUM2, QTYF, SEC_FILL, TITLE, _cell, _num, _setup_print


def _title(ws, title, project_name, label, ncols):
    ws.cell(1, 1, title).font = TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(1, 1).alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28
    ws.cell(2, 1, f"공사명 : {project_name}").font = FONT
    ws.cell(2, max(ncols - 2, 2), label).font = BOLD
    ws.cell(2, ncols, "(단위 : 원)").font = FONT


def write_change_workbook(path: Path, project, dc, lines: list[ChangeLine], summary: dict) -> Path:
    wb = openpyxl.Workbook()
    label = f"{dc.change_no}차 설계변경"

    # ---- 1. 설계변경 대비표 (좌우 3블록)
    ws = wb.active
    ws.title = "설계변경대비표"
    heads = [("품명", 24), ("규격", 22), ("단위", 6),
             ("당초\n수량", 10), ("당초\n단가", 11), ("당초\n금액", 13),
             ("변경\n수량", 10), ("변경\n단가", 11), ("변경\n금액", 13),
             ("증감\n수량", 10), ("증감\n단가", 11), ("증감\n금액", 13), ("구분", 8), ("비고", 14)]
    n = len(heads)
    _title(ws, "설 계 변 경 대 비 표", project.name, label, n)
    for i, (h, w) in enumerate(heads, start=1):
        _cell(ws, 4, i, h, bold=True, align="center", fill=HEAD_FILL)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 30
    r = 5
    sec_rows = []
    cur = None
    first = None
    for cl in lines:
        if cl.row_type == "section":
            if cur is not None and first is not None and r > first:
                sec_rows.append((cur, first, r - 1))
            _cell(ws, r, 1, " > ".join(cl.section_path) or cl.name, bold=True, fill=SEC_FILL)
            for c in range(2, n + 1):
                _cell(ws, r, c, fill=SEC_FILL)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            cur, first = r, r + 1
            r += 1
            continue
        o, c_, d = cl.orig, cl.chg, cl.diff
        _cell(ws, r, 1, cl.name)
        _cell(ws, r, 2, cl.spec)
        _cell(ws, r, 3, cl.unit, align="center")
        _cell(ws, r, 4, _num(o.get("qty")), QTYF)
        _cell(ws, r, 5, _num(o.get("tot_price")), NUM)
        _cell(ws, r, 6, _num(o.get("tot_amt")), NUM)
        _cell(ws, r, 7, _num(c_.get("qty")), QTYF)
        _cell(ws, r, 8, _num(c_.get("tot_price")), NUM)
        _cell(ws, r, 9, _num(c_.get("tot_amt")), NUM)
        _cell(ws, r, 10, f"=G{r}-D{r}", QTYF)
        _cell(ws, r, 11, f"=H{r}-E{r}", NUM)
        _cell(ws, r, 12, f"=I{r}-F{r}", NUM)
        _cell(ws, r, 13, {"add": "신규", "delete": "삭제", "modify": "변경", "keep": ""}[cl.action], align="center")
        _cell(ws, r, 14, cl.price_basis[:40] if cl.action in ("add", "modify") and cl.price_type != "contract" else (cl.item.note if cl.item else ""))
        r += 1
    if cur is not None and first is not None and r > first:
        sec_rows.append((cur, first, r - 1))
    for sr, a, b in sec_rows:
        for c in (6, 9, 12):
            col = get_column_letter(c)
            ws.cell(sr, c, f"=SUM({col}{a}:{col}{b})").number_format = NUM
            ws.cell(sr, c).font = BOLD
    _cell(ws, r, 1, "총    계", bold=True, align="center", fill=HEAD_FILL)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for c in range(2, n + 1):
        _cell(ws, r, c, fill=HEAD_FILL)
    for c in (6, 9, 12):
        col = get_column_letter(c)
        refs = ",".join(f"{col}{sr}" for sr, _, _ in sec_rows) or f"{col}5:{col}{r-1}"
        ws.cell(r, c, f"=SUM({refs})").number_format = NUM
        ws.cell(r, c).font = BOLD
    total_row = r
    ws.freeze_panes = "D5"
    _setup_print(ws, f"{project.name} {label}")

    # ---- 2. 변경내역서 (당초/변경 2행형, 재료·노무·경비 분리)
    ws2 = wb.create_sheet("변경내역서")
    h2 = [("품명", 24), ("규격", 22), ("단위", 6), ("구분", 6), ("수량", 10), ("재료비\n단가", 10), ("재료비\n금액", 12),
          ("노무비\n단가", 10), ("노무비\n금액", 12), ("경비\n단가", 9), ("경비\n금액", 11), ("합계\n단가", 11), ("합계\n금액", 13), ("증감금액", 13), ("비고", 12)]
    n2 = len(h2)
    _title(ws2, "설 계 변 경 내 역 서", project.name, label, n2)
    for i, (h, w) in enumerate(h2, start=1):
        _cell(ws2, 4, i, h, bold=True, align="center", fill=HEAD_FILL)
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.row_dimensions[4].height = 30
    r = 5
    for cl in lines:
        if cl.row_type == "section":
            _cell(ws2, r, 1, " > ".join(cl.section_path) or cl.name, bold=True, fill=SEC_FILL)
            for c in range(2, n2 + 1):
                _cell(ws2, r, c, fill=SEC_FILL)
            ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            r += 1
            continue
        for tag, dd in (("당초", cl.orig), ("변경", cl.chg)):
            _cell(ws2, r, 1, cl.name if tag == "당초" else "")
            _cell(ws2, r, 2, cl.spec if tag == "당초" else "")
            _cell(ws2, r, 3, cl.unit if tag == "당초" else "", align="center")
            _cell(ws2, r, 4, tag, align="center")
            _cell(ws2, r, 5, _num(dd.get("qty")), QTYF)
            _cell(ws2, r, 6, _num(dd.get("mat_price")), NUM)
            _cell(ws2, r, 7, _num(dd.get("mat_amt")), NUM)
            _cell(ws2, r, 8, _num(dd.get("lab_price")), NUM)
            _cell(ws2, r, 9, _num(dd.get("lab_amt")), NUM)
            _cell(ws2, r, 10, _num(dd.get("exp_price")), NUM)
            _cell(ws2, r, 11, _num(dd.get("exp_amt")), NUM)
            _cell(ws2, r, 12, f"=F{r}+H{r}+J{r}", NUM)
            _cell(ws2, r, 13, f"=G{r}+I{r}+K{r}", NUM)
            if tag == "당초":
                _cell(ws2, r, 14, f"=M{r+1}-M{r}", NUM)
                _cell(ws2, r, 15, {"add": "신규", "delete": "삭제", "modify": "", "keep": ""}[cl.action] or (cl.item.note if cl.item else ""))
            else:
                _cell(ws2, r, 14)
                _cell(ws2, r, 15, cl.price_basis[:30] if cl.price_type != "contract" else "")
            r += 1
    ws2.freeze_panes = "E5"
    _setup_print(ws2, f"{project.name} {label}")

    # ---- 3. 증감 총괄표 (공종별)
    ws3 = wb.create_sheet("증감총괄표")
    h3 = [("공종", 44), ("당초금액", 16), ("변경금액", 16), ("증감금액", 16)]
    _title(ws3, "공 사 비 증 · 감 총 괄 표", project.name, label, 4)
    for i, (h, w) in enumerate(h3, start=1):
        _cell(ws3, 4, i, h, bold=True, align="center", fill=HEAD_FILL)
        ws3.column_dimensions[get_column_letter(i)].width = w
    rr = 5
    for sr, a, b in sec_rows:
        _cell(ws3, rr, 1, ws.cell(sr, 1).value)
        _cell(ws3, rr, 2, f"=설계변경대비표!F{sr}", NUM)
        _cell(ws3, rr, 3, f"=설계변경대비표!I{sr}", NUM)
        _cell(ws3, rr, 4, f"=C{rr}-B{rr}", NUM)
        rr += 1
    _cell(ws3, rr, 1, "직접공사비 계", bold=True, fill=HEAD_FILL)
    for c in (2, 3, 4):
        col = get_column_letter(c)
        _cell(ws3, rr, c, f"=SUM({col}5:{col}{rr-1})", NUM, bold=True, fill=HEAD_FILL)
    _setup_print(ws3, f"{project.name} {label}", landscape=False)

    # ---- 4. 변경 원가계산서
    if summary.get("cost_sheet"):
        ws4 = wb.create_sheet("변경원가계산서")
        h4 = [("비목", 30), ("당초 금액", 16), ("변경 금액", 16), ("증감", 16), ("요율(%)", 9), ("산정기준", 24), ("비고", 10)]
        _title(ws4, "변 경 원 가 계 산 서", project.name, label, 7)
        for i, (h, w) in enumerate(h4, start=1):
            _cell(ws4, 4, i, h, bold=True, align="center", fill=HEAD_FILL)
            ws4.column_dimensions[get_column_letter(i)].width = w
        rr = 5
        for cs in summary["cost_sheet"]:
            bold = cs["kind"] in ("subtotal", "total")
            _cell(ws4, rr, 1, cs["name"], bold=bold, fill=SEC_FILL if bold else None)
            _cell(ws4, rr, 2, _num(Decimal(cs["orig"])) if cs["orig"] is not None else None, NUM, bold=bold)
            _cell(ws4, rr, 3, _num(Decimal(cs["chg"])) if cs["chg"] is not None else None, NUM, bold=bold)
            _cell(ws4, rr, 4, f"=C{rr}-B{rr}", NUM, bold=bold)
            _cell(ws4, rr, 5, float(cs["rate"]) if cs["rate"] else None, "0.###")
            _cell(ws4, rr, 6, cs["base_text"] or cs["rate_base"])
            _cell(ws4, rr, 7, "정산" if cs["fixed"] else "")
            rr += 1
        _setup_print(ws4, f"{project.name} {label}", landscape=False)

    # ---- 5. 신규 품목 내역서
    news = [cl for cl in lines if cl.action == "add" and cl.row_type == "item"]
    if news:
        ws5 = wb.create_sheet("신규품목내역서")
        h5 = [("공종", 28), ("품명", 24), ("규격", 22), ("단위", 6), ("수량", 10), ("재료비단가", 11), ("노무비단가", 11), ("경비단가", 10), ("합계단가", 11), ("금액", 13), ("단가유형", 14), ("산정근거", 30)]
        _title(ws5, "신 규 품 목 내 역 서", project.name, label, len(h5))
        for i, (h, w) in enumerate(h5, start=1):
            _cell(ws5, 4, i, h, bold=True, align="center", fill=HEAD_FILL)
            ws5.column_dimensions[get_column_letter(i)].width = w
        rr = 5
        from ..services.design_change import PRICE_TYPES
        for cl in news:
            _cell(ws5, rr, 1, " > ".join(cl.section_path))
            _cell(ws5, rr, 2, cl.name)
            _cell(ws5, rr, 3, cl.spec)
            _cell(ws5, rr, 4, cl.unit, align="center")
            _cell(ws5, rr, 5, _num(cl.chg["qty"]), QTYF)
            _cell(ws5, rr, 6, _num(cl.chg["mat_price"]), NUM)
            _cell(ws5, rr, 7, _num(cl.chg["lab_price"]), NUM)
            _cell(ws5, rr, 8, _num(cl.chg["exp_price"]), NUM)
            _cell(ws5, rr, 9, f"=F{rr}+G{rr}+H{rr}", NUM)
            _cell(ws5, rr, 10, _num(cl.chg["tot_amt"]), NUM)
            _cell(ws5, rr, 11, PRICE_TYPES.get(cl.price_type, cl.price_type))
            _cell(ws5, rr, 12, cl.price_basis)
            rr += 1
        _setup_print(ws5, f"{project.name} {label}")

    # ---- 6. 변경 사유
    ws6 = wb.create_sheet("설계변경사유")
    ws6.cell(1, 1, "설 계 변 경 사 유 서").font = TITLE
    ws6.cell(3, 1, f"공사명 : {project.name}").font = FONT
    ws6.cell(4, 1, f"변경차수 : {label}").font = FONT
    ws6.cell(5, 1, f"적용일 : {dc.applied_date or ''}").font = FONT
    ws6.cell(7, 1, "변경 사유").font = BOLD
    ws6.cell(8, 1, dc.reason or "").alignment = Alignment(wrap_text=True, vertical="top")
    ws6.column_dimensions["A"].width = 100
    ws6.row_dimensions[8].height = 200
    wb.save(path)
    return path
