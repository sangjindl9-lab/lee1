"""시트 역할 자동 분류 (이름 사전 + 헤더 키워드 + 구조 마커)."""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl.worksheet.worksheet import Worksheet

from .loader import SheetSummary, cell_text, norm_text

# 역할 코드
COST = "cost_sheet"          # 원가계산서
SUMMARY = "summary"          # 총괄표/집계표
ESTIMATE = "estimate"        # 공종별 내역서
CHANGE_ESTIMATE = "change_estimate"  # 변경내역서(당초/변경 2행)
CHANGE_COMPARE = "change_compare"    # 증감대비표
UNIT_PRICE = "unit_price"    # 일위대가
UNIT_PRICE_INDEX = "unit_price_index"  # 일위대가 목차
MATERIAL = "material_price"  # 자재단가/합산자재/단가조사/대비표
LABOR = "labor_rate"         # 노임
QUANTITY = "quantity"        # 수량/공량 산출
REFERENCE = "reference"      # 표지/목차/옵션/시방서 등
UNKNOWN = "unknown"

ROLE_LABELS = {
    COST: "원가계산서", SUMMARY: "총괄표/집계표", ESTIMATE: "공종별 내역서",
    CHANGE_ESTIMATE: "변경내역서", CHANGE_COMPARE: "증감대비표", UNIT_PRICE: "일위대가표",
    UNIT_PRICE_INDEX: "일위대가 목록", MATERIAL: "자재단가/단가조사", LABOR: "노임단가",
    QUANTITY: "수량/공량 산출서", REFERENCE: "참조/기타", UNKNOWN: "미확인",
}

NAME_RULES: list[tuple[str, list[str]]] = [
    (MATERIAL, ["단가대비", "자재단가", "단가비교", "합산자재", "단가조사"]),
    (CHANGE_COMPARE, ["증감대비", "증감내역", "증감총괄", "대비표"]),
    (CHANGE_ESTIMATE, ["변경내역"]),
    (COST, ["원가계산", "원가", "공사비예산", "공사비총괄", "갑지"]),
    (UNIT_PRICE_INDEX, ["일대목차", "일위대가목록", "일위목록", "일위대가목차"]),
    (LABOR, ["노임", "일위노임", "노임근거", "노무산출", "노무비산출"]),
    (UNIT_PRICE, ["일위대가", "일위", "호표"]),
    (MATERIAL, ["합산자재", "단가조사", "단가대비", "자재단가", "단가비교", "자재집계", "단가산출", "공구손료", "기계경비"]),
    (QUANTITY, ["산출집계", "기본산출", "공량산출", "수량산출", "자재산출", "산출근거", "산출서"]),
    (SUMMARY, ["총괄표", "총괄내역", "집계표", "공종별집계", "공종리스트"]),
    (ESTIMATE, ["내역서", "내역", "설계명세", "산출내역", "명세서"]),
    (REFERENCE, ["표지", "목차", "옵션", "시방", "보안", "예정", "동원", "결재", "사유서", "도면", "계통도",
                 "공사계획", "공사설명", "자금계획", "용량산출", "케이블list", "sheet", "사용설명", "한전불입", "적용근거"]),
]

HEADER_KEYS = {
    "name": ["품명", "명칭", "공종명", "품목"],
    "spec": ["규격"],
    "unit": ["단위"],
    "qty": ["수량"],
    "mat": ["재료비"],
    "lab": ["노무비"],
    "exp": ["경비"],
    "amt": ["금액"],
    "price": ["단가"],
    "orig": ["당초", "변경전"],
    "chg": ["변경", "변경후"],
    "diff": ["증감", "증", "감"],
    "rate": ["요율", "비율", "구성비"],
    "item": ["비목", "구분"],
    "labor": ["직종", "노임단가", "할증"],
}


@dataclass
class SheetRole:
    sheet: str
    role: str
    confidence: float
    reasons: list[str]
    header_row: int | None = None

    def as_dict(self):
        return {"sheet": self.sheet, "role": self.role, "label": ROLE_LABELS.get(self.role, self.role),
                "confidence": round(self.confidence, 2), "reasons": self.reasons, "header_row": self.header_row}


def _scan_header(ws: Worksheet, max_rows: int = 12) -> tuple[set[str], int | None, set[str]]:
    """상단 행에서 헤더 키워드 집합과 헤더 행 번호, 구조 마커를 찾는다."""
    found: set[str] = set()
    header_row = None
    best = 0
    markers: set[str] = set()
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows), start=1):
        keys_in_row: set[str] = set()
        for c in row:
            t = norm_text(c.value)
            if not t:
                continue
            for key, words in HEADER_KEYS.items():
                for w in words:
                    if w in t and len(t) <= 12:
                        keys_in_row.add(key)
        if len(keys_in_row) > best and ("name" in keys_in_row or "item" in keys_in_row):
            best = len(keys_in_row)
            header_row = r_idx
        found |= keys_in_row
    # 구조 마커 (B열)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400), min_col=1, max_col=3):
        for c in row:
            t = cell_text(c.value)
            if t in ("공종줄", "합계줄", "일목줄", "소계줄"):
                markers.add(t)
    return found, header_row, markers


def classify_sheet(ws: Worksheet, summary: SheetSummary) -> SheetRole:
    nm = norm_text(ws.title)
    reasons: list[str] = []
    role = UNKNOWN
    conf = 0.0

    # 1) 이름 사전
    for r, words in NAME_RULES:
        if any(w in nm for w in words):
            role, conf = r, 0.6
            reasons.append(f"시트명 '{ws.title}' 사전 일치")
            break

    # 2) 헤더/구조 확인
    keys, header_row, markers = _scan_header(ws)
    has_item_cols = {"name", "unit", "qty"} <= keys or ({"name", "qty"} <= keys and "amt" in keys)
    cost_cols = {"mat", "lab", "exp"} & keys

    if "일목줄" in markers:
        if role not in (UNIT_PRICE,):
            role = UNIT_PRICE
            reasons.append("B열 '일목줄' 마커")
        conf = max(conf, 0.95)
    elif "공종줄" in markers and has_item_cols:
        if role in (ESTIMATE, SUMMARY, CHANGE_ESTIMATE):
            conf = max(conf, 0.95)
        elif role in (LABOR, MATERIAL, QUANTITY, UNIT_PRICE_INDEX, COST):
            conf = max(conf, 0.8)   # 이름이 강한 근거 - 유지
        else:
            role = ESTIMATE
            conf = max(conf, 0.85)
            reasons.append("B열 '공종줄' 마커 + 품명/수량 열")
        reasons.append("공종줄/합계줄 구조 마커 확인")
    elif has_item_cols and {"orig", "chg"} <= keys and "diff" in keys:
        role = CHANGE_COMPARE if role in (UNKNOWN, CHANGE_COMPARE, ESTIMATE) else role
        conf = max(conf, 0.85)
        reasons.append("당초/변경/증감 열 구조")
    elif has_item_cols and role in (UNKNOWN, REFERENCE):
        role = ESTIMATE
        conf = max(conf, 0.65)
        reasons.append("품명/단위/수량/금액 헤더")
    elif has_item_cols:
        conf = max(conf, 0.8)
        reasons.append("품명/단위/수량 헤더 확인")
    elif "item" in keys and "rate" in keys and ("amt" in keys or "orig" in keys):
        if role in (UNKNOWN, REFERENCE, SUMMARY):
            role = COST
        conf = max(conf, 0.85)
        reasons.append("비목/요율/금액 헤더 (원가계산서 구조)")
    elif "labor" in keys and role in (UNKNOWN, UNIT_PRICE, LABOR):
        conf = max(conf, 0.75)
        reasons.append("직종/노임단가 헤더")
    elif role == UNKNOWN:
        if summary.value_count < 5:
            role, conf = REFERENCE, 0.5
            reasons.append("데이터가 거의 없음")
        else:
            conf = 0.3
            reasons.append("헤더 패턴 미확인 - 사용자 지정 필요")
    else:
        conf = max(conf, 0.5)

    if summary.hidden and role in (REFERENCE, UNKNOWN):
        reasons.append("숨김 시트")
    return SheetRole(sheet=ws.title, role=role, confidence=conf, reasons=reasons, header_row=header_row)


def classify_workbook(bundle) -> list[SheetRole]:
    out = []
    for s in bundle.sheets:
        ws = bundle.formulas[s.name]
        out.append(classify_sheet(ws, s))
    # 내역서가 여러 개면 가장 큰 것을 주 내역서로 (뒤에서 선택)
    return out
