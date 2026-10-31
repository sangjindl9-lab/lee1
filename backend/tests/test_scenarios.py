"""필수 검증 시나리오 (요구사항 19항). 순서대로 실행되며 상태를 공유한다."""
import io
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.util import today_kst
from conftest import ok

S: dict = {}
T = today_kst()
D1, D2 = (T - timedelta(days=2)).isoformat(), (T - timedelta(days=1)).isoformat()


@pytest.fixture
def admin(client_factory):
    return client_factory().login("admin", "Admin1234!")


def test_01_admin_creates_users_and_assigns_site(admin, client_factory):
    roles = {r["code"]: r["id"] for r in ok(admin.get("/api/admin/roles"))["roles"]}
    S["roles"] = roles
    staff = ok(admin.post("/api/admin/users", json={"login_id": "staff1", "name": "현장직원1", "role_id": roles["site_staff"],
                                                    "password": "Staff1234!"}))
    mgr = ok(admin.post("/api/admin/users", json={"login_id": "mgr1", "name": "현장소장1", "role_id": roles["site_manager"],
                                                  "password": "Mgr12345!"}))
    buyer = ok(admin.post("/api/admin/users", json={"login_id": "buyer1", "name": "구매1", "role_id": roles["purchase"],
                                                    "password": "Buy12345!"}))
    S.update(staff=staff["id"], mgr=mgr["id"], buyer=buyer["id"])
    # 중복 아이디 거부
    assert admin.post("/api/admin/users", json={"login_id": "staff1", "name": "x", "role_id": roles["site_staff"],
                                                "password": "Staff1234!"}).status_code == 409
    p1 = ok(admin.post("/api/projects", json={"code": "P-001", "name": "A현장 전기공사", "status": "진행",
                                              "supply_amount": "100000000", "start_date": D1, "site_manager_id": mgr["id"]}))
    p2 = ok(admin.post("/api/projects", json={"code": "P-002", "name": "B현장 통신공사", "status": "진행"}))
    assert p1["vat_amount"] == "10000000" and p1["total_amount"] == "110000000"
    S.update(p1=p1["id"], p2=p2["id"])
    ok(admin.put(f"/api/projects/{p1['id']}/members", json=[{"user_id": staff["id"], "role_in_project": "공사"}]))
    # 최초 로그인: 비밀번호 변경 전 업무 API 차단
    c = client_factory().login("staff1", "Staff1234!")
    assert c.get("/api/projects").status_code == 428
    ok(c.post("/api/auth/password", json={"current_password": "Staff1234!", "new_password": "Staff5678!"}))
    for login, old, new in (("mgr1", "Mgr12345!", "Mgr56789!"), ("buyer1", "Buy12345!", "Buy56789!")):
        cc = client_factory().login(login, old)
        ok(cc.post("/api/auth/password", json={"current_password": old, "new_password": new}))


def test_02_staff_only_sees_assigned_site(client_factory):
    c = client_factory().login("staff1", "Staff5678!")
    ids = [p["id"] for p in ok(c.get("/api/projects"))]
    assert ids == [S["p1"]]
    assert c.get(f"/api/projects/{S['p2']}").status_code == 403
    assert c.get(f"/api/projects/{S['p2']}/boq").status_code == 403
    assert c.get(f"/api/daily-reports?project_id={S['p2']}").json() == []
    # 민감정보(계약금액) 비노출
    assert ok(c.get(f"/api/projects/{S['p1']}"))["supply_amount"] is None


def test_03_progress_1000m_200_100(admin, client_factory):
    g = ok(admin.post(f"/api/projects/{S['p1']}/boq/groups", json={"name": "케이블 포설"}))
    it = ok(admin.post(f"/api/projects/{S['p1']}/boq/items", json={
        "group_id": g["id"], "name": "F-CV 케이블", "spec": "4C 16SQ", "unit": "m", "design_qty": "1000",
        "mat_unit_price": "1000", "labor_unit_price": "500"}))
    assert it["amount"] == "1500000"
    zero = ok(admin.post(f"/api/projects/{S['p1']}/boq/items", json={"group_id": g["id"], "name": "수량0 항목", "unit": "식",
                                                                     "design_qty": "0", "mat_unit_price": "100"}))
    S.update(item=it["id"], zero=zero["id"])
    staff = client_factory().login("staff1", "Staff5678!")
    mgr = client_factory().login("mgr1", "Mgr56789!")
    r1 = ok(staff.post("/api/daily-reports", json={"project_id": S["p1"], "report_date": D1, "work_summary": "케이블 포설",
                                                   "items": [{"boq_item_id": it["id"], "qty_today": "200"}],
                                                   "labor": [{"trade": "내선전공", "headcount": "4"}]}))
    ok(staff.post(f"/api/daily-reports/{r1['id']}/submit"))
    # 현장 직원은 승인 불가
    assert staff.post(f"/api/daily-reports/{r1['id']}/approve").status_code == 403
    ok(mgr.post(f"/api/daily-reports/{r1['id']}/approve"))
    r2 = ok(staff.post("/api/daily-reports", json={"project_id": S["p1"], "report_date": D2, "work_summary": "케이블 포설 계속",
                                                   "items": [{"boq_item_id": it["id"], "qty_today": "100"}]}))
    ok(staff.post(f"/api/daily-reports/{r2['id']}/submit"))
    d = ok(mgr.post(f"/api/daily-reports/{r2['id']}/approve"))
    row = d["items"][0]
    assert (row["prev_cum"], row["qty_today"], row["cum"], row["remaining"], row["rate"]) == ("200", "100", "300", "700", "30")
    pr = ok(admin.get(f"/api/projects/{S['p1']}/boq/progress"))
    item = next(i for i in pr["items"] if i["id"] == it["id"])
    assert item["cum_qty"] == "300" and item["remaining_qty"] == "700" and item["rate"] == "30"
    zi = next(i for i in pr["items"] if i["id"] == zero["id"])
    assert zi["rate"] is None and zi["excluded_reason"].startswith("설계수량 0")
    assert pr["actual_rate"] == "30"  # 금액가중, 수량0 항목 제외
    S.update(r1=r1["id"], r2=r2["id"])


def test_04_resave_no_duplicate(admin, client_factory):
    staff = client_factory().login("staff1", "Staff5678!")
    # 승인된 일보는 직접 수정 불가
    r = staff.put(f"/api/daily-reports/{S['r2']}", json={"work_summary": "x", "items": [{"boq_item_id": S["item"], "qty_today": "100"}]})
    assert r.status_code == 409
    # 같은 일보 재승인 시도 → 상태 오류, 누계 그대로
    mgr = client_factory().login("mgr1", "Mgr56789!")
    assert mgr.post(f"/api/daily-reports/{S['r2']}/approve").status_code == 409
    # 작성중 일보를 여러 번 저장해도 실적 미반영
    r3 = ok(staff.post("/api/daily-reports", json={"project_id": S["p1"], "report_date": T.isoformat(), "work_summary": "임시",
                                                   "items": [{"boq_item_id": S["item"], "qty_today": "50"}]}))
    for _ in range(3):
        r3 = ok(staff.put(f"/api/daily-reports/{r3['id']}", json={"version": r3["version"], "work_summary": "임시",
                                                                    "items": [{"boq_item_id": S["item"], "qty_today": "50"}]}))
    # 오래된 version 으로 저장 → 충돌
    assert staff.put(f"/api/daily-reports/{r3['id']}", json={"version": 1, "work_summary": "x", "items": []}).status_code == 409
    # 같은 날짜 일보 중복 생성 불가
    assert staff.post("/api/daily-reports", json={"project_id": S["p1"], "report_date": T.isoformat(), "work_summary": "중복"}).status_code == 409
    pr = ok(admin.get(f"/api/projects/{S['p1']}/boq/progress"))
    assert next(i for i in pr["items"] if i["id"] == S["item"])["cum_qty"] == "300"
    S["r3"] = r3["id"]


def test_05_correction_and_unapprove(admin, client_factory):
    staff = client_factory().login("staff1", "Staff5678!")
    mgr = client_factory().login("mgr1", "Mgr56789!")
    ok(staff.post(f"/api/daily-reports/{S['r2']}/request-correction", json={"reason": "수량 오기 정정"}))
    # 정정 중에는 기존 승인 실적 유지
    pr = ok(admin.get(f"/api/projects/{S['p1']}/boq/progress"))
    assert next(i for i in pr["items"] if i["id"] == S["item"])["cum_qty"] == "300"
    d = ok(staff.get(f"/api/daily-reports/{S['r2']}"))
    ok(staff.put(f"/api/daily-reports/{S['r2']}", json={"version": d["version"], "work_summary": "케이블 포설 계속(정정)",
                                                          "items": [{"boq_item_id": S["item"], "qty_today": "150"}]}))
    ok(staff.post(f"/api/daily-reports/{S['r2']}/submit"))
    d = ok(mgr.post(f"/api/daily-reports/{S['r2']}/approve"))
    assert d["revision"] == 2 and d["items"][0]["cum"] == "350"
    pr = ok(admin.get(f"/api/projects/{S['p1']}/boq/progress"))
    assert next(i for i in pr["items"] if i["id"] == S["item"])["cum_qty"] == "350"
    assert len(d["revisions"]) == 2  # 1차 승인본 원문 보존
    old = ok(admin.get(f"/api/daily-reports/{S['r2']}/revisions/1"))
    assert old["items"][0]["qty_today"] == "100"
    # 승인취소 → 실적 제거
    assert staff.post(f"/api/daily-reports/{S['r2']}/unapprove", json={"reason": "x"}).status_code == 403
    ok(admin.post(f"/api/daily-reports/{S['r2']}/unapprove", json={"reason": "현장 재확인 필요"}))
    pr = ok(admin.get(f"/api/projects/{S['p1']}/boq/progress"))
    it = next(i for i in pr["items"] if i["id"] == S["item"])
    assert it["cum_qty"] == "200" and it["rate"] == "20"
    # 1일차 승인본의 고정 수치는 변하지 않음, 설계수량 변경 후에도 과거 승인본 유지
    b = ok(admin.get(f"/api/projects/{S['p1']}/boq"))
    cur = next(i for i in b["items"] if i["id"] == S["item"])
    ok(admin.put(f"/api/projects/{S['p1']}/boq/items/{S['item']}", json={**{k: cur[k] for k in (
        "group_id", "name", "spec", "unit", "mat_unit_price", "labor_unit_price", "exp_unit_price")},
        "design_qty": "1200", "version": cur["version"]}))
    r1 = ok(admin.get(f"/api/daily-reports/{S['r1']}"))
    assert r1["items"][0]["design_qty"] == "1000" and r1["items"][0]["snapshot"] is True


def test_06_material_request_approval_po_partial_receipt(admin, client_factory):
    staff = client_factory().login("staff1", "Staff5678!")
    mgr = client_factory().login("mgr1", "Mgr56789!")
    buyer = client_factory().login("buyer1", "Buy56789!")
    v = ok(buyer.post("/api/vendors", json={"name": "테스트전선상사"}))
    mr = ok(staff.post("/api/material-requests", json={"project_id": S["p1"], "need_date": T.isoformat(), "reason": "간선 포설용",
                                                       "items": [{"name": "F-CV", "spec": "4C 16SQ", "unit": "m", "qty": "100"}]}))
    ok(staff.post(f"/api/material-requests/{mr['id']}/submit", json={"steps": [{"user_id": S["mgr"], "kind": "결재"}]}))
    # 결재중 수정 불가
    assert staff.put(f"/api/material-requests/{mr['id']}", json={"project_id": S["p1"], "items": [
        {"name": "F-CV", "unit": "m", "qty": "200"}]}).status_code == 409
    # 결재 차례가 아닌 사람은 승인 불가
    doc_id = ok(staff.get(f"/api/material-requests/{mr['id']}"))["approval"]["id"]
    assert buyer.post(f"/api/approvals/{doc_id}/act", json={"action": "approve"}).status_code == 403
    ok(mgr.post(f"/api/approvals/{doc_id}/act", json={"action": "approve", "comment": "승인"}))
    mr = ok(staff.get(f"/api/material-requests/{mr['id']}"))
    assert mr["status"] == "approved"
    ri = mr["items"][0]["id"]
    # 초과 발주 차단
    bad = buyer.post("/api/purchase-orders", json={"project_id": S["p1"], "vendor_id": v["id"], "request_id": mr["id"],
                                                   "items": [{"request_item_id": ri, "name": "F-CV", "unit": "m", "qty": "120", "unit_price": "9500"}]})
    assert bad.status_code == 400
    po = ok(buyer.post("/api/purchase-orders", json={"project_id": S["p1"], "vendor_id": v["id"], "request_id": mr["id"],
                                                     "delivery_date": D1,
                                                     "items": [{"request_item_id": ri, "name": "F-CV", "unit": "m", "qty": "100", "unit_price": "9500"}]}))
    assert po["supply_amount"] == "950000" and po["vat_amount"] == "95000" and po["total_amount"] == "1045000"
    # 확정 전 입고 불가
    key = uuid.uuid4().hex
    assert staff.post(f"/api/purchase-orders/{po['id']}/receipts", json={"request_key": key, "items": [
        {"po_item_id": po["items"][0]["id"], "qty": "10"}]}).status_code == 409
    ok(buyer.post(f"/api/purchase-orders/{po['id']}/submit", json={"steps": [{"user_id": S["mgr"], "kind": "결재"}]}))
    po = ok(buyer.get(f"/api/purchase-orders/{po['id']}"))
    ok(mgr.post(f"/api/approvals/{po['approval']['id']}/act", json={"action": "approve"}))
    po = ok(buyer.get(f"/api/purchase-orders/{po['id']}"))
    assert po["status"] == "ordered"
    pi = po["items"][0]["id"]
    # 60개 입고 → 미입고 40
    po = ok(staff.post(f"/api/purchase-orders/{po['id']}/receipts", json={"request_key": key, "items": [{"po_item_id": pi, "qty": "60"}]}))
    assert po["items"][0]["outstanding_qty"] == "40" and po["status"] == "partial"
    # 같은 요청 키 재전송(중복 클릭) → 중복 입고 없음
    po = ok(staff.post(f"/api/purchase-orders/{po['id']}/receipts", json={"request_key": key, "items": [{"po_item_id": pi, "qty": "60"}]}))
    assert po["duplicate"] is True and po["items"][0]["outstanding_qty"] == "40"
    # 미입고 초과 입고 차단
    assert staff.post(f"/api/purchase-orders/{po['id']}/receipts", json={"request_key": uuid.uuid4().hex,
                      "items": [{"po_item_id": pi, "qty": "41"}]}).status_code == 400
    # 반품 10 → 순입고 50, 미입고 50
    po = ok(staff.post(f"/api/purchase-orders/{po['id']}/receipts", json={"kind": "반품", "request_key": uuid.uuid4().hex,
                                                                         "items": [{"po_item_id": pi, "qty": "10"}]}))
    it = po["items"][0]
    assert (it["received_qty"], it["returned_qty"], it["net_received_qty"], it["outstanding_qty"]) == ("60", "10", "50", "50")
    # 입고분이 있으면 전체취소 불가, 잔량취소 가능
    assert buyer.post(f"/api/purchase-orders/{po['id']}/cancel", json={"mode": "all", "reason": "x"}).status_code == 409
    po = ok(buyer.post(f"/api/purchase-orders/{po['id']}/cancel", json={"mode": "remaining", "reason": "현장 사용량 감소"}))
    assert po["items"][0]["cancelled_qty"] == "50" and po["items"][0]["outstanding_qty"] == "0" and po["status"] == "received"
    # 현장 직원은 단가 비노출
    assert ok(staff.get(f"/api/purchase-orders/{po['id']}"))["items"][0]["unit_price"] is None
    # 외부 발송은 연동 전 차단
    assert buyer.post(f"/api/purchase-orders/{po['id']}/send").status_code == 501
    S.update(po=po["id"], vendor=v["id"])


def test_07_rejected_doc_not_treated_as_approved(client_factory):
    staff = client_factory().login("staff1", "Staff5678!")
    mgr = client_factory().login("mgr1", "Mgr56789!")
    buyer = client_factory().login("buyer1", "Buy56789!")
    mr = ok(staff.post("/api/material-requests", json={"project_id": S["p1"], "items": [{"name": "UTP", "unit": "m", "qty": "50"}]}))
    ok(staff.post(f"/api/material-requests/{mr['id']}/submit", json={"steps": [{"user_id": S["mgr"], "kind": "결재"}]}))
    doc = ok(staff.get(f"/api/material-requests/{mr['id']}"))["approval"]["id"]
    assert mgr.post(f"/api/approvals/{doc}/act", json={"action": "reject"}).status_code == 400  # 사유 필수
    ok(mgr.post(f"/api/approvals/{doc}/act", json={"action": "reject", "comment": "수량 근거 부족"}))
    mr2 = ok(staff.get(f"/api/material-requests/{mr['id']}"))
    assert mr2["status"] == "rejected"
    assert buyer.post("/api/purchase-orders", json={"project_id": S["p1"], "vendor_id": S["vendor"], "request_id": mr["id"],
                      "items": [{"request_item_id": mr2["items"][0]["id"], "name": "UTP", "unit": "m", "qty": "50"}]}).status_code == 409
    assert buyer.post(f"/api/material-requests/{mr['id']}/quotes", json={"vendor_id": S["vendor"], "items": []}).status_code == 409
    # 반려된 문서를 다시 승인 처리 시도 → 불가
    assert mgr.post(f"/api/approvals/{doc}/act", json={"action": "approve"}).status_code == 409
    # 수정 후 재상신 가능
    ok(staff.put(f"/api/material-requests/{mr['id']}", json={"project_id": S["p1"], "reason": "근거 추가",
                                                              "items": [{"name": "UTP", "unit": "m", "qty": "40"}]}))
    ok(staff.post(f"/api/material-requests/{mr['id']}/submit", json={"steps": [{"user_id": S["mgr"], "kind": "결재"}]}))
    d = ok(mgr.get(f"/api/approvals/{doc}"))
    assert d["status"] == "in_progress" and d["round"] == 2
    # 회수
    ok(staff.post(f"/api/approvals/{doc}/withdraw"))
    assert ok(staff.get(f"/api/material-requests/{mr['id']}"))["status"] == "draft"


def test_08_access_control_server_side(client_factory):
    from fastapi.testclient import TestClient
    from app.main import app
    anon = TestClient(app)
    assert anon.get("/api/projects").status_code == 401
    staff = client_factory().login("staff1", "Staff5678!")
    for url in ("/api/admin/users", "/api/admin/roles", "/api/admin/audit", "/api/admin/trash"):
        assert staff.get(url).status_code == 403, url
    roles = S["roles"]
    assert staff.post("/api/admin/users", json={"login_id": "hack", "name": "h", "role_id": roles["admin"],
                                                "password": "Hack1234!"}).status_code == 403
    assert staff.put("/api/admin/company", json={"name": "x"}).status_code == 403
    assert staff.post("/api/projects", json={"code": "X", "name": "x"}).status_code == 403
    # CSRF 토큰 없는 변경 요청 차단
    r = staff.c.post("/api/projects", json={"code": "X", "name": "x"})
    assert r.status_code == 403
    # 비활성화된 사용자는 즉시 차단
    admin = client_factory().login("admin", "Admin1234!")
    u = next(x for x in ok(admin.get("/api/admin/users")) if x["login_id"] == "buyer1")
    ok(admin.put(f"/api/admin/users/{u['id']}", json={"name": u["name"], "role_id": u["role_id"], "is_active": False}))
    assert client_factory().c.post("/api/auth/login", json={"login_id": "buyer1", "password": "Buy56789!"}).status_code == 403
    ok(admin.put(f"/api/admin/users/{u['id']}", json={"name": u["name"], "role_id": u["role_id"], "is_active": True}))
    # 비밀번호 초기화
    tmp = ok(admin.post(f"/api/admin/users/{u['id']}/reset-password"))["temporary_password"]
    b = client_factory().login("buyer1", tmp)
    assert b.get("/api/projects").status_code == 428
    ok(b.post("/api/auth/password", json={"current_password": tmp, "new_password": "Buy56789!"}))
    # 변경 이력 기록 확인
    logs = ok(admin.get("/api/admin/audit?entity=daily_report"))
    assert any(l["action"] == "unapprove" for l in logs)


def _png():
    buf = io.BytesIO()
    Image.new("RGB", (800, 600), (40, 80, 160)).save(buf, "PNG")
    return buf.getvalue()


def test_09_files_persist_across_sessions(client_factory):
    staff = client_factory().login("staff1", "Staff5678!")
    r = staff.post("/api/files", data={"owner_type": "daily_report", "owner_id": str(S["r3"]), "category": "사진", "phase": "작업중",
                                       "location": "1층 EPS", "title": "케이블 포설"},
                   files=[("files", ("현장사진.png", _png(), "image/png")), ("files", ("사진2.png", _png(), "image/png"))])
    up = ok(r)
    assert len(up) == 2 and up[0]["has_thumb"]
    # 실행파일 차단
    bad = staff.post("/api/files", data={"owner_type": "project", "owner_id": str(S["p1"])},
                     files=[("files", ("a.exe", b"MZ...", "application/octet-stream"))])
    assert bad.status_code == 400
    # 다른 현장에 업로드 불가
    assert staff.post("/api/files", data={"owner_type": "project", "owner_id": str(S["p2"])},
                      files=[("files", ("a.pdf", b"%PDF-1.4", "application/pdf"))]).status_code == 403
    # 새 세션(재로그인)에서 조회·다운로드
    again = client_factory().login("staff1", "Staff5678!")
    lst = ok(again.get(f"/api/files?owner_type=daily_report&owner_id={S['r3']}"))
    assert len(lst) == 2
    dl = again.get(f"/api/files/{lst[0]['id']}/download")
    assert dl.status_code == 200 and dl.content[:4] == b"\x89PNG"
    S["photo_ids"] = [x["id"] for x in lst]


def test_10_pdf_and_excel(admin, client_factory, tmp_path):
    import subprocess
    outs = {
        "report.pdf": f"/api/export/daily-reports/{S['r1']}.pdf",
        "po.pdf": f"/api/export/purchase-orders/{S['po']}.pdf",
        "progress.pdf": f"/api/export/projects/{S['p1']}/progress.pdf",
        "ledger.pdf": f"/api/export/projects/{S['p1']}/photo-ledger.pdf?per_page=4",
    }
    for name, url in outs.items():
        r = admin.get(url)
        assert r.status_code == 200, (url, r.text[:300])
        assert r.content[:4] == b"%PDF"
        p = tmp_path / name
        p.write_bytes(r.content)
        text = subprocess.run(["pdftotext", str(p), "-"], capture_output=True, text=True).stdout
        assert any(k in text for k in ("작 업 일 보", "발 주 서", "공 정 현 황", "사 진 대 장")), name
        assert "A현장" in text or "케이블" in text
    for url, sheet in ((f"/api/export/daily-reports.xlsx?project_id={S['p1']}", "시공수량"),
                       (f"/api/export/purchase-orders/{S['po']}.xlsx", "발주서"),
                       (f"/api/export/projects/{S['p1']}/boq.xlsx", "설계내역·공정"),
                       ("/api/import/template/boq.xlsx", "업로드양식")):
        r = admin.get(url)
        assert r.status_code == 200, url
        wb = load_workbook(io.BytesIO(r.content))
        assert sheet in wb.sheetnames
    S["pdf_dir"] = str(tmp_path)


def test_11_excel_import_preview_validate_commit(admin):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["대공종", "중공종", "품명", "규격", "단위", "설계수량", "재료비단가", "노무비단가", "경비단가", "계획시작일", "계획종료일"])
    ws.append(["접지", "", "접지봉", "14x1000", "본", "20", "15000", "8000", "0", "2026-11-01", "2026-11-10"])
    ws.append(["접지", "", "", "", "본", "abc", "", "", "", "", ""])          # 품명 누락 + 숫자오류
    ws.append(["접지", "", "접지봉", "14x1000", "본", "5", "", "", "", "", ""])  # 파일 내 중복
    buf = io.BytesIO()
    wb.save(buf)
    r = ok(admin.post("/api/import/preview", data={"kind": "boq", "project_id": str(S["p1"])},
                      files={"file": ("내역.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}))
    valid = [x["index"] for x in r["rows"] if x["valid"]]
    assert valid == [0] and len(r["rows"]) == 3
    assert any("품명 누락" in e for e in r["rows"][1]["errors"]) and any("파일 내 중복" in e for e in r["rows"][2]["errors"])
    assert admin.post(f"/api/import/{r['job_id']}/commit", json={"mapping": r["mapping"], "selected": [0, 1]}).status_code == 400
    res = ok(admin.post(f"/api/import/{r['job_id']}/commit", json={"mapping": r["mapping"], "selected": [0]}))
    assert res["inserted"] == 1
    b = ok(admin.get(f"/api/projects/{S['p1']}/boq"))
    assert any(i["name"] == "접지봉" and i["amount"] == "460000" for i in b["items"])


def test_12_dashboard(admin, client_factory):
    d = ok(admin.get("/api/dashboard"))
    assert d["counts"]["active"] >= 2
    p = next(x for x in d["progress"] if x["id"] == S["p1"])
    assert p["actual_rate"] is not None
    staff = client_factory().login("staff1", "Staff5678!")
    ds = ok(staff.get("/api/dashboard"))
    assert all(x["id"] == S["p1"] for x in ds["progress"]) and ds["money"] is None
