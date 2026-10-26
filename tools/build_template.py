"""원본 구매요청서 엑셀에서 '내장 양식'(빈 양식)을 만든다.
- 출력되는 4개 시트(구매요청서·구매내역현황·관급단가 비교표·단가및인증현황조사)의
  글꼴·테두리·채우기·맞춤·표시형식·열너비·행높이·병합·인증마크 그림·인쇄설정을 그대로 복사
- 다른 공사 데이터, 숨김 시트, 외부 연결은 버림
- 품목 CAP개까지 들어가도록 줄/블록을 늘림
"""
import copy, io, sys
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.worksheet.pagebreak import Break, RowBreak

SRC = sys.argv[1] if len(sys.argv) > 1 else 'template.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'gwangeup_template.xlsx'
CAP = 20   # 품목 최대 개수

src = openpyxl.load_workbook(SRC)
S_REQ, S_STA = src['구매요청서'], src['구매내역현황']
S_CMP, S_SUR = src['관급단가 비교표_케이블트레이'], src['단가및인증현황조사(변경)_케이블트레이']

wb = openpyxl.Workbook(); wb.remove(wb.active)
N_REQ, N_STA, N_CMP, N_SUR = '구매요청서', '구매내역현황', '관급단가 비교표', '단가및인증현황조사'
CMP_REF, SUR_REF = f"'{N_CMP}'", N_SUR


from openpyxl.cell.cell import MergedCell
def clear(cell):
    if not isinstance(cell, MergedCell):
        cell.value = None


def copy_cell(sc, dc, value=True):
    if isinstance(dc, MergedCell):
        return
    if value:
        dc.value = None if isinstance(sc, MergedCell) else sc.value
    if sc.has_style:
        dc.font = copy.copy(sc.font); dc.border = copy.copy(sc.border); dc.fill = copy.copy(sc.fill)
        dc.number_format = sc.number_format; dc.protection = copy.copy(sc.protection); dc.alignment = copy.copy(sc.alignment)


def copy_rows(s, d, src_rows, dst_start, max_col, value=True):
    for i, r in enumerate(src_rows):
        dr = dst_start + i
        if r in s.row_dimensions and s.row_dimensions[r].height:
            d.row_dimensions[dr].height = s.row_dimensions[r].height
        for c in range(1, max_col + 1):
            copy_cell(s.cell(r, c), d.cell(dr, c), value)


def copy_cols(s, d, max_col):
    for k, cd in s.column_dimensions.items():
        if cd.min and cd.min > max_col:
            continue
        for c in range(cd.min or CI(k), min(cd.max or CI(k), max_col) + 1):
            x = d.column_dimensions[L(c)]
            x.width = cd.width; x.hidden = cd.hidden


def copy_merges(s, d, r1, r2, max_col, offset):
    for m in s.merged_cells.ranges:
        if m.min_row >= r1 and m.max_row <= r2 and m.max_col <= max_col:
            d.merge_cells(start_row=m.min_row + offset, end_row=m.max_row + offset, start_column=m.min_col, end_column=m.max_col)


def copy_page(s, d, area):
    d.page_setup.orientation = s.page_setup.orientation
    d.page_setup.paperSize = s.page_setup.paperSize
    d.page_setup.scale = s.page_setup.scale
    for a in ('left', 'right', 'top', 'bottom', 'header', 'footer'):
        setattr(d.page_margins, a, getattr(s.page_margins, a))
    d.print_options.horizontalCentered = s.print_options.horizontalCentered
    d.print_options.verticalCentered = s.print_options.verticalCentered
    d.sheet_view.showGridLines = s.sheet_view.showGridLines
    d.sheet_view.zoomScale = s.sheet_view.zoomScale
    d.print_area = area


_BYTES = {}


def clone_image(img, d, col, row, colOff, rowOff, tcol, trow, tcolOff, trowOff):
    if id(img) not in _BYTES:
        _BYTES[id(img)] = img._data()
    data = _BYTES[id(img)]
    import hashlib, os
    os.makedirs('_imgcache', exist_ok=True)
    p = os.path.join('_imgcache', hashlib.md5(data).hexdigest() + '.' + (img.format or 'png'))
    if not os.path.exists(p):
        open(p, 'wb').write(data)
    ni = XLImage(p)
    ni.anchor = TwoCellAnchor(_from=AnchorMarker(col=col, colOff=colOff, row=row, rowOff=rowOff),
                              to=AnchorMarker(col=tcol, colOff=tcolOff, row=trow, rowOff=trowOff), editAs='oneCell')
    d.add_image(ni)


# ───────────────────── 1) 구매요청서 ─────────────────────
d = wb.create_sheet(N_REQ)
copy_cols(S_REQ, d, 14)
copy_rows(S_REQ, d, range(1, 45), 1, 14)
copy_merges(S_REQ, d, 1, 15, 14, 0)
for r in range(16, 45):  # 원본 22~24행의 불규칙 병합 대신 품목줄과 같은 병합으로 통일
    d.merge_cells(f'A{r}:E{r}'); d.merge_cells(f'F{r}:I{r}')
# 원본 수식/값 정리 후 내장 연결식
for r in range(13, 45):
    for c in range(1, 15):
        clear(d.cell(r, c))
d['A13'] = 'Ⅱ 전기공사 관급자재'; d['M13'] = '=SUM(M15)'
d['A14'] = '1. 도급자관급'; d['M14'] = '=M15'
d['A15'] = ''; d['M15'] = f'=SUM(M16:M{15 + CAP})'
d['N13'] = ''
for k in range(CAP):
    r, sr = 16 + k, 10 + k
    d[f'A{r}'] = f'=IF({N_STA}!A{sr}="","",{N_STA}!A{sr})'
    d[f'F{r}'] = f'=IF({N_STA}!A{sr}="","",{N_STA}!B{sr})'
    d[f'J{r}'] = f'=IF({N_STA}!A{sr}="","",{N_STA}!C{sr})'
    d[f'K{r}'] = f'=IF({N_STA}!A{sr}="","",{N_STA}!N{sr})'
    d[f'L{r}'] = f'=IF({N_STA}!A{sr}="","",{N_STA}!O{sr})'
    d[f'M{r}'] = f'=IF({N_STA}!A{sr}="","",TRUNC(K{r}*L{r},0))'
    # 행 스타일: 16행(품목줄) 기준
    for c in range(1, 15):
        copy_cell(S_REQ.cell(16, c), d.cell(r, c), value=False)
    d.row_dimensions[r].height = S_REQ.row_dimensions[16].height
for r in range(16 + CAP, 45):
    for c in range(1, 15):
        copy_cell(S_REQ.cell(24, c), d.cell(r, c), value=False)
copy_page(S_REQ, d, 'A1:N44')
d.print_title_rows = '12:12'

# ───────────────────── 2) 구매내역현황 ─────────────────────
d = wb.create_sheet(N_STA)
MC = 20  # A~T
copy_cols(S_STA, d, MC)
copy_rows(S_STA, d, range(1, 8), 1, MC)
copy_merges(S_STA, d, 1, 6, MC, 0)
for c in range(1, MC + 1):
    clear(d.cell(7, c))
copy_rows(S_STA, d, [8], 8, MC, value=False)
copy_rows(S_STA, d, [39], 9, MC, value=False)
for k in range(CAP):
    copy_rows(S_STA, d, [40], 10 + k, MC, value=False)
last = 9 + CAP
d['A7'] = ' Ⅱ. 관급자재대'; d['A8'] = '1. 도급자관급'; d['A9'] = ''
for col in 'IJMPS':
    d[f'{col}7'] = f'={col}8'; d[f'{col}8'] = f'={col}9'; d[f'{col}9'] = f'=SUM({col}10:{col}{last})'
for k in range(CAP):
    r = 10 + k
    d[f'I{r}'] = f'=IF(A{r}="","",E{r}*H{r})'
    d[f'J{r}'] = f'=IF(A{r}="","",G{r}*H{r})'
    d[f'M{r}'] = f'=IF(A{r}="","",TRUNC(K{r}*L{r}))'
    d[f'O{r}'] = f'=IF(A{r}="","",{SUR_REF}!BF{6 + k})'
    d[f'P{r}'] = f'=IF(A{r}="","",TRUNC(N{r}*O{r}))'
    d[f'Q{r}'] = f'=IF(A{r}="","",E{r}-K{r}-N{r})'
    d[f'R{r}'] = f'=IF(A{r}="","",H{r}-O{r})'
    d[f'S{r}'] = f'=IF(A{r}="","",J{r}-M{r}-P{r})'
copy_page(S_STA, d, f'A1:T{last}')
d.print_title_rows = '5:6'

# ───────────────────── 3) 관급단가 비교표 ─────────────────────
d = wb.create_sheet(N_CMP)
copy_cols(S_CMP, d, 17)
icons = [im for im in S_CMP._images if im.anchor._from.row < 13 and im.anchor._from.col != 15]
for k in range(CAP):
    off = 13 * k
    copy_rows(S_CMP, d, range(1, 14), 1 + off, 17)
    copy_merges(S_CMP, d, 1, 13, 17, off)
    for r in range(1 + off, 14 + off):
        for c in (1, 2):
            clear(d.cell(r, c))
        for c in range(4, 18):
            if 5 + off <= r <= 11 + off:
                clear(d.cell(r, c))
    d.cell(1 + off, 3).value = ' □ 품목명 : '
    for im in icons:
        f, t = im.anchor._from, im.anchor.to
        clone_image(im, d, f.col, f.row + off, f.colOff, f.rowOff, t.col, t.row + off, t.colOff, t.rowOff)
    if k:
        d.row_breaks.append(Break(id=off))
copy_page(S_CMP, d, 'C1:Q13')

# ───────────────────── 4) 단가및인증현황조사 ─────────────────────
d = wb.create_sheet(N_SUR)
MC = 58  # A~BF
copy_cols(S_SUR, d, MC)
copy_rows(S_SUR, d, range(1, 6), 1, MC)
copy_merges(S_SUR, d, 1, 5, MC, 0)
tot = 6 + CAP
for k in range(CAP):
    copy_rows(S_SUR, d, [6], 6 + k, MC, value=False)
copy_rows(S_SUR, d, [31], tot, MC, value=False)
for m in S_SUR.merged_cells.ranges:
    if m.min_row == 31 and m.max_col <= MC:
        d.merge_cells(start_row=tot, end_row=tot, start_column=m.min_col, end_column=m.max_col)
d['A1'] = '관급자재 조달단가 및 인증현황 조사표'
d['H3'] = f"={CMP_REF}!D5"; d['S3'] = f"={CMP_REF}!D6"; d['AB3'] = f"={CMP_REF}!D7"
for c in ('AL3', 'AV3'):
    d[c] = None
for k in range(CAP):
    r, b = 6 + k, 5 + 13 * k
    d[f'D{r}'] = f'=IF(A{r}="","",{N_STA}!N{10 + k})'
    d[f'G{r}'] = f'=IF(A{r}="","",D{r})'
    for (no, pr, am), j in ((('H', 'I', 'J'), 0), (('R', 'S', 'T'), 1), (('AB', 'AC', 'AD'), 2)):
        d[f'{no}{r}'] = f'=IF({CMP_REF}!F{b + j}="","",{CMP_REF}!F{b + j})'
        d[f'{pr}{r}'] = f'=IF({CMP_REF}!H{b + j}="","",{CMP_REF}!H{b + j})'
        d[f'{am}{r}'] = f'=IF(OR(A{r}="",{pr}{r}=""),"",TRUNC($G{r}*{pr}{r}))'
    d[f'BF{r}'] = f'=IF(A{r}="","",TRUNC(AVERAGE(S{r},AC{r}),0))'
end = tot - 1
d[f'A{tot}'] = '합    계'
d[f'J{tot}'] = f'=SUM(J6:J{end})'; d[f'K{tot}'] = f'=COUNTA(K6:Q{end})'
d[f'T{tot}'] = f'=SUM(T6:T{end})'; d[f'V{tot}'] = f'=COUNTA(U6:AA{end})'
d[f'AD{tot}'] = f'=SUM(AD6:AD{end})'; d[f'AE{tot}'] = f'=COUNTA(AE6:AK{end})'
for im in S_SUR._images:
    f, t = im.anchor._from, im.anchor.to
    if f.col < MC and not S_SUR.column_dimensions[L(f.col + 1)].hidden and f.row == 4:
        clone_image(im, d, f.col, f.row, f.colOff, f.rowOff, t.col, t.row, t.colOff, t.rowOff)
copy_page(S_SUR, d, f'A1:BF{tot}')
d.row_breaks.append(Break(id=tot + 1))

wb.active = 0
wb.save(OUT)
print('saved', OUT, 'CAP', CAP)
