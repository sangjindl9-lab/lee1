"""원본 엑셀 로더.

- 수식본(formulas)과 값본(data_only)을 동시에 읽는다.
- 외부 연결, 오류 셀, 매크로, 숨김 시트, 시트 간 참조 관계를 수집한다.
- 원본 파일은 절대 수정하지 않는다 (읽기 전용 로드).
"""
from __future__ import annotations

import re
import warnings
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

warnings.filterwarnings("ignore", module="openpyxl")

SHEET_REF_RE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_가-힣\.\-\(\) ]+))!\$?[A-Z]{1,3}\$?\d+")
EXTERNAL_RE = re.compile(r"\[\d+\]")
ERROR_VALUES = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}


@dataclass
class SheetSummary:
    name: str
    index: int
    state: str
    max_row: int
    max_col: int
    used_range: str
    formula_count: int = 0
    value_count: int = 0
    merged_count: int = 0
    references_to: dict[str, int] = field(default_factory=dict)  # 참조 대상 시트 -> 횟수
    external_ref_count: int = 0
    error_cells: list[str] = field(default_factory=list)
    hidden: bool = False


@dataclass
class WorkbookBundle:
    path: Path
    formulas: openpyxl.Workbook
    values: openpyxl.Workbook
    sheets: list[SheetSummary]
    has_macro: bool = False
    is_xls_converted: bool = False
    warnings: list[str] = field(default_factory=list)

    def sheet_names(self) -> list[str]:
        return [s.name for s in self.sheets]

    def summary(self) -> dict[str, Any]:
        return {
            "file_name": self.path.name,
            "sheet_count": len(self.sheets),
            "has_macro": self.has_macro,
            "is_xls_converted": self.is_xls_converted,
            "warnings": self.warnings,
            "sheets": [
                {
                    "name": s.name,
                    "index": s.index,
                    "state": s.state,
                    "hidden": s.hidden,
                    "used_range": s.used_range,
                    "max_row": s.max_row,
                    "max_col": s.max_col,
                    "formula_count": s.formula_count,
                    "value_count": s.value_count,
                    "merged_count": s.merged_count,
                    "references_to": s.references_to,
                    "external_ref_count": s.external_ref_count,
                    "error_cells": s.error_cells[:50],
                    "error_count": len(s.error_cells),
                }
                for s in self.sheets
            ],
        }


def _is_encrypted(path: Path) -> bool:
    try:
        with open(path, "rb") as fh:
            head = fh.read(8)
        # OLE compound file (xls or encrypted ooxml)
        if head.startswith(b"\xd0\xcf\x11\xe0") and path.suffix.lower() != ".xls":
            return True
    except OSError:
        pass
    return False


def _convert_xls(path: Path, out_dir: Path) -> tuple[Path, list[str]]:
    """구형 .xls → .xlsx 중간 파일 (값만 보존, 수식은 유실됨)."""
    import xlrd

    warns = ["구형 .xls 파일은 값만 변환됩니다. 수식과 서식 일부가 유실될 수 있습니다."]
    book = xlrd.open_workbook(path, formatting_info=False)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():
        ws = wb.create_sheet(sh.name[:31])
        if sh.visibility:
            ws.sheet_state = "hidden"
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if v not in ("", None):
                    ws.cell(r + 1, c + 1, v)
    out = out_dir / (path.stem + "_converted.xlsx")
    wb.save(out)
    return out, warns


def load_bundle(path: str | Path, work_dir: str | Path | None = None) -> WorkbookBundle:
    path = Path(path)
    warns: list[str] = []
    converted = False
    if _is_encrypted(path):
        raise ValueError("암호가 설정된 파일은 처리할 수 없습니다. 암호를 해제한 뒤 다시 업로드하세요.")
    load_path = path
    if path.suffix.lower() == ".xls":
        out_dir = Path(work_dir) if work_dir else path.parent
        load_path, w = _convert_xls(path, out_dir)
        warns += w
        converted = True

    has_macro = False
    if load_path.suffix.lower() == ".xlsm":
        try:
            with zipfile.ZipFile(load_path) as z:
                has_macro = any(n.lower().endswith("vbaproject.bin") for n in z.namelist())
        except zipfile.BadZipFile:
            pass
        if has_macro:
            warns.append("매크로(VBA)가 포함된 파일입니다. 매크로는 실행되지 않으며 결과물은 .xlsx로 생성됩니다.")

    wb_f = openpyxl.load_workbook(load_path, data_only=False, read_only=False, keep_links=False)
    wb_v = openpyxl.load_workbook(load_path, data_only=True, read_only=False, keep_links=False)

    sheets: list[SheetSummary] = []
    for idx, ws in enumerate(wb_f.worksheets):
        s = SheetSummary(
            name=ws.title,
            index=idx,
            state=ws.sheet_state,
            hidden=ws.sheet_state != "visible",
            max_row=ws.max_row,
            max_col=ws.max_column,
            used_range=ws.dimensions,
            merged_count=len(ws.merged_cells.ranges),
        )
        wsv = wb_v[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None or v == "":
                    continue
                if isinstance(v, str) and v.startswith("="):
                    s.formula_count += 1
                    if EXTERNAL_RE.search(v):
                        s.external_ref_count += 1
                    for m in SHEET_REF_RE.finditer(v):
                        target = (m.group(1) or m.group(2) or "").strip()
                        if target and target != ws.title:
                            s.references_to[target] = s.references_to.get(target, 0) + 1
                    cv = wsv[cell.coordinate].value
                    if isinstance(cv, str) and cv in ERROR_VALUES:
                        s.error_cells.append(f"{cell.coordinate}:{cv}")
                else:
                    s.value_count += 1
                    if isinstance(v, str) and v in ERROR_VALUES:
                        s.error_cells.append(f"{cell.coordinate}:{v}")
        if s.external_ref_count:
            warns.append(f"[{ws.title}] 외부 파일 참조 수식 {s.external_ref_count}개 - 저장된 값만 사용합니다.")
        if s.error_cells:
            warns.append(f"[{ws.title}] 오류 셀 {len(s.error_cells)}개 (#REF!, #DIV/0! 등) - 검증 목록에 표시됩니다.")
        sheets.append(s)

    return WorkbookBundle(path=path, formulas=wb_f, values=wb_v, sheets=sheets, has_macro=has_macro,
                          is_xls_converted=converted, warnings=warns)


def cell_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).replace("\n", " ").strip()


def norm_text(v: Any) -> str:
    """헤더 비교용: 공백/구두점 제거."""
    return re.sub(r"[\s\.\,\(\)\[\]_\-:：·]", "", cell_text(v)).lower()


def iter_sheet(ws: Worksheet, max_rows: int | None = None):
    for i, row in enumerate(ws.iter_rows(values_only=False), start=1):
        if max_rows and i > max_rows:
            break
        yield i, row
