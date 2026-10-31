"""엑셀(.xlsx) 내보내기·업로드 검증."""
import io
from datetime import date, datetime
from decimal import Decimal

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="DDE4F0")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _cell_value(v):
    if isinstance(v, Decimal):
        return float(v) if v != v.to_integral_value() else int(v)
    return v


def add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int] | None = None,
              number_cols: set[int] | None = None, caption: str | None = None):
    ws = wb.create_sheet(title[:31])
    r0 = 1
    if caption:
        ws.cell(1, 1, caption).font = Font(bold=True, size=12)
        r0 = 3
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(r0, c, h)
        cell.font, cell.fill, cell.border = Font(bold=True), HEAD_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ri, row in enumerate(rows, start=r0 + 1):
        for c, v in enumerate(row, start=1):
            if isinstance(v, str) and number_cols and (c - 1) in number_cols:
                try:
                    v = Decimal(v)
                except Exception:
                    pass
            cell = ws.cell(ri, c, _cell_value(v))
            cell.border = BORDER
            if number_cols and (c - 1) in number_cols:
                cell.number_format = "#,##0.###"
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths[c - 1] if widths and c - 1 < len(widths) else 14)
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{max(r0, r0 + len(rows))}"
    ws.print_options.gridLines = False
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = f"{r0}:{r0}"
    return ws


def workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- 업로드 양식/검증
BOQ_FIELDS = [  # key, 헤더명, 필수
    ("group1", "대공종", True), ("group2", "중공종", False), ("name", "품명", True), ("spec", "규격", False),
    ("unit", "단위", True), ("design_qty", "설계수량", True), ("mat_unit_price", "재료비단가", False),
    ("labor_unit_price", "노무비단가", False), ("exp_unit_price", "경비단가", False),
    ("plan_start", "계획시작일", False), ("plan_end", "계획종료일", False), ("custom_weight", "가중치", False),
]
MATERIAL_FIELDS = [
    ("code", "자재코드", True), ("name", "품명", True), ("spec", "규격", False), ("unit", "단위", True),
    ("maker", "제조사", False), ("model", "모델명", False), ("ref_price", "참고단가", False),
    ("stock_managed", "재고관리(Y/N)", False), ("min_stock", "최소재고", False), ("supply_type", "관급/사급", False),
]
COMMON_UNITS = {"m", "km", "M", "EA", "ea", "개", "식", "대", "본", "면", "조", "SET", "set", "롤", "kg", "ton", "매",
                "개소", "회선", "구", "포", "Lot", "LOT", "m2", "㎡", "m3", "㎥", "SPAN", "경간", "주", "인", "일", "시간", "공"}


def fields_for(kind: str):
    if kind == "boq":
        return BOQ_FIELDS
    if kind == "material":
        return MATERIAL_FIELDS
    raise HTTPException(400, "알 수 없는 업로드 종류")


def template_bytes(kind: str) -> bytes:
    f = fields_for(kind)
    wb = workbook()
    sample = {
        "boq": [["배관", "전선관", "HI-PVC 전선관", "16C", "m", 1000, 1200, 3500, 100, "2026-11-01", "2026-11-30", ""],
                ["배선", "", "F-GV", "2.5SQ", "m", 3000, 450, 900, 0, "2026-11-15", "2026-12-15", ""]],
        "material": [["E-0001", "HI-PVC 전선관", "16C", "m", "", "", 1200, "N", 0, "사급"]],
    }[kind]
    ws = add_sheet(wb, "업로드양식", [h + (" *" if req else "") for _, h, req in f], sample, widths=[16] * len(f))
    note = wb.create_sheet("작성안내")
    note["A1"] = "* 표시 열은 필수입니다. 예시 행은 지우고 작성하세요."
    note["A2"] = "날짜는 YYYY-MM-DD 형식, 금액·수량은 숫자만 입력하세요(쉼표 가능)."
    note["A3"] = "업로드 후 미리보기 화면에서 열 매핑과 오류를 확인하고, 정상 행만 선택해 반영합니다."
    return to_bytes(wb)


def read_sheet(data: bytes) -> tuple[list[str], list[list]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "엑셀 파일(.xlsx)을 읽을 수 없습니다. 다른 이름으로 저장 → Excel 통합 문서로 저장 후 다시 시도하세요.")
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else (v.date().isoformat() if isinstance(v, datetime) else
                                          v.isoformat() if isinstance(v, date) else str(v).strip()) for v in row])
    while rows and not any(rows[0]):
        rows.pop(0)
    if not rows:
        raise HTTPException(400, "빈 시트입니다.")
    headers = [h.replace("*", "").strip() for h in rows[0]]
    body = [r for r in rows[1:] if any(c != "" for c in r)]
    if len(body) > 5000:
        raise HTTPException(400, "한 번에 5,000행까지 업로드할 수 있습니다.")
    return headers, body


def auto_mapping(kind: str, headers: list[str]) -> dict[str, int | None]:
    m = {}
    norm = [h.replace(" ", "").lower() for h in headers]
    for key, label, _ in fields_for(kind):
        lab = label.replace(" ", "").lower()
        idx = next((i for i, h in enumerate(norm) if h == lab), None)
        if idx is None:
            idx = next((i for i, h in enumerate(norm) if h and (lab.startswith(h) or h.startswith(lab.split("(")[0]))), None)
        m[key] = idx
    return m
