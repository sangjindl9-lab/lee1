"""원본 파일 업로드·분석·매핑·가져오기·미리보기."""
from __future__ import annotations

import hashlib
import shutil
import uuid
from pathlib import Path

import openpyxl
from fastapi import APIRouter, BackgroundTasks, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session

from ..auth import audit, get_current_user, get_project_for
from ..config import settings
from ..db import SessionLocal, get_db
from ..excel.loader import cell_text
from ..models import MappingRule, SourceFile, User, WorkbookSheet
from ..schemas import ImportIn, SheetUpdate
from ..services.importer import import_contract, run_analysis

router = APIRouter(prefix="/api")
ALLOWED = {".xlsx", ".xlsm", ".xls"}


def _file_out(f: SourceFile) -> dict:
    a = f.analysis or {}
    return {
        "id": f.id, "project_id": f.project_id, "name": f.original_name, "size": f.size, "purpose": f.purpose,
        "status": f.status, "error": f.error, "uploaded_at": f.uploaded_at.isoformat(), "sha256": f.sha256,
        "summary": {
            "sheet_count": a.get("file", {}).get("sheet_count"),
            "has_macro": a.get("file", {}).get("has_macro"),
            "warnings": a.get("file", {}).get("warnings", []),
            "estimate_sheets": a.get("estimate_sheets", []),
            "primary_cost_sheet": a.get("primary_cost_sheet"),
            "has_cost_sheet": a.get("has_cost_sheet"),
            "has_unit_price": a.get("has_unit_price"),
            "grand_totals": a.get("grand_totals", {}),
            "unidentified": a.get("unidentified", []),
        },
        "sheets": [{"id": s.id, "name": s.name, "index": s.index, "role": s.role, "user_role": s.user_role,
                    "confidence": s.role_confidence, "hidden": s.hidden, "has_map": bool(s.column_map),
                    "item_count": (s.stats or {}).get("item_count"), "finding_count": (s.stats or {}).get("finding_count")}
                   for s in sorted(f.sheets, key=lambda x: x.index)],
    }


def _analyze_bg(file_id: int):
    db = SessionLocal()
    try:
        sf = db.get(SourceFile, file_id)
        if sf:
            try:
                run_analysis(db, sf, str(settings.storage_dir / "work"))
            except Exception:  # noqa: BLE001
                pass
    finally:
        db.close()


@router.post("/projects/{pid}/files")
async def upload(pid: int, bg: BackgroundTasks, file: UploadFile = File(...), purpose: str = Form("contract"),
                 sync: bool = Form(False), db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    p = get_project_for(db, pid, user, "edit")
    ext = Path(file.filename or "").suffix.lower()
    if ext not in ALLOWED:
        raise HTTPException(400, "엑셀 파일(.xlsx, .xlsm, .xls)만 업로드할 수 있습니다.")
    pdir = settings.storage_dir / f"project_{p.id}"
    pdir.mkdir(parents=True, exist_ok=True)
    stored = pdir / f"{uuid.uuid4().hex}{ext}"
    h = hashlib.sha256()
    size = 0
    with open(stored, "wb") as out:
        while chunk := await file.read(1024 * 1024):
            size += len(chunk)
            if size > settings.max_upload_mb * 1024 * 1024:
                out.close()
                stored.unlink(missing_ok=True)
                raise HTTPException(413, f"파일이 {settings.max_upload_mb}MB를 초과합니다.")
            h.update(chunk)
            out.write(chunk)
    stored.chmod(0o444)  # 원본 읽기 전용 보관
    sf = SourceFile(project_id=p.id, original_name=file.filename or stored.name, stored_path=str(stored), size=size,
                    sha256=h.hexdigest(), purpose=purpose, uploaded_by=user.id)
    db.add(sf)
    audit(db, user, "file.upload", p.id, target=sf.original_name)
    db.commit()
    if purpose == "contract":
        if sync:
            try:
                run_analysis(db, sf, str(settings.storage_dir / "work"))
            except Exception as e:  # noqa: BLE001
                raise HTTPException(400, f"분석 실패: {e}")
        else:
            bg.add_task(_analyze_bg, sf.id)
    return _file_out(sf)


@router.get("/projects/{pid}/files")
def list_files(pid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    get_project_for(db, pid, user)
    return [_file_out(f) for f in db.query(SourceFile).filter_by(project_id=pid).order_by(SourceFile.id.desc())]


def _get_file(db, fid, user, need="view") -> SourceFile:
    sf = db.get(SourceFile, fid)
    if not sf:
        raise HTTPException(404, "파일을 찾을 수 없습니다.")
    get_project_for(db, sf.project_id, user, need)
    return sf


@router.get("/files/{fid}")
def get_file(fid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return _file_out(_get_file(db, fid, user))


@router.get("/files/{fid}/analysis")
def get_analysis(fid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    sf = _get_file(db, fid, user)
    if sf.status != "analyzed":
        return {"status": sf.status, "error": sf.error}
    a = dict(sf.analysis)
    a["status"] = "analyzed"
    return a


@router.post("/files/{fid}/analyze")
def reanalyze(fid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    sf = _get_file(db, fid, user, "edit")
    try:
        run_analysis(db, sf, str(settings.storage_dir / "work"))
    except Exception as e:  # noqa: BLE001
        raise HTTPException(400, f"분석 실패: {e}")
    return _file_out(sf)


@router.get("/files/{fid}/download")
def download(fid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    sf = _get_file(db, fid, user)
    return FileResponse(sf.stored_path, filename=sf.original_name)


@router.delete("/files/{fid}")
def delete_file(fid: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """원본 파일 삭제: 이 파일에서 가져온 계약내역(버전)·원가비목·단가표·일위대가도 함께 삭제.
    그 계약내역을 쓰는 기성 회차가 있으면 삭제하지 않고 안내."""
    from ..models import ContractVersion, PriceTable, UnitPriceAnalysis
    from .estimates import delete_version_rows
    sf = _get_file(db, fid, user, "edit")
    versions = db.query(ContractVersion).filter_by(source_file_id=sf.id).order_by(ContractVersion.version_no.desc()).all()
    ids = {v.id for v in versions}
    for v in versions:
        delete_version_rows(db, v, ids)   # 기성 회차가 있으면 400 → 아무것도 지우지 않음
    db.query(PriceTable).filter_by(source_file_id=sf.id).delete()
    db.query(UnitPriceAnalysis).filter_by(source_file_id=sf.id).delete()
    audit(db, user, "file.delete", sf.project_id, target=sf.original_name, after={"deleted_versions": len(versions)})
    db.delete(sf)
    db.commit()
    return {"ok": True, "deleted_versions": len(versions)}


@router.get("/files/{fid}/preview")
def preview(fid: int, sheet: str, start: int = 1, rows: int = 60, cols: int = 30,
            db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """원본 시트 원시 셀 값 (열 매핑 편집기용)."""
    sf = _get_file(db, fid, user)
    wb = openpyxl.load_workbook(sf.stored_path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise HTTPException(404, "시트가 없습니다.")
    ws = wb[sheet]
    out = []
    for r_idx, row in enumerate(ws.iter_rows(min_row=start, max_row=start + rows - 1, max_col=cols, values_only=True), start=start):
        out.append([cell_text(v)[:60] for v in row])
    wb.close()
    return {"sheet": sheet, "start": start, "rows": out, "max_row": ws.max_row, "max_col": ws.max_column}


@router.put("/files/{fid}/sheets/{sheet_name}")
def update_sheet(fid: int, sheet_name: str, body: SheetUpdate, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """시트 역할 / 열 매핑 사용자 지정 → 재분석."""
    sf = _get_file(db, fid, user, "edit")
    ws = next((s for s in sf.sheets if s.name == sheet_name), None)
    if not ws:
        raise HTTPException(404, "시트가 없습니다.")
    if body.user_role is not None:
        ws.user_role = body.user_role or None
    if body.user_column_map is not None:
        ws.user_column_map = body.user_column_map
    if body.save_as_rule and ws.user_column_map:
        sig = hashlib.sha1("|".join(f"{k}:{v}" for k, v in sorted((ws.column_map.get("raw_headers") or {}).items())).encode()).hexdigest()[:16]
        rule = db.query(MappingRule).filter_by(signature=sig).first()
        if not rule:
            rule = MappingRule(signature=sig, created_by=user.id)
            db.add(rule)
        rule.name = f"{sf.original_name}/{sheet_name}"
        rule.role = ws.user_role or ws.role
        rule.column_map = ws.user_column_map
    audit(db, user, "sheet.mapping", sf.project_id, target=f"{sf.original_name}/{sheet_name}",
          after={"user_role": ws.user_role, "user_column_map": ws.user_column_map})
    db.commit()
    run_analysis(db, sf, str(settings.storage_dir / "work"))
    return _file_out(sf)


@router.post("/files/{fid}/import")
def import_file(fid: int, body: ImportIn, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """분석된 내역서를 계약내역 버전으로 가져오기."""
    sf = _get_file(db, fid, user, "edit")
    p = get_project_for(db, sf.project_id, user, "edit")
    if sf.status != "analyzed":
        raise HTTPException(400, "먼저 분석이 완료되어야 합니다.")
    role_over = {ws.name: ws.user_role for ws in sf.sheets if ws.user_role}
    map_over = {ws.name: ws.user_column_map for ws in sf.sheets if ws.user_column_map}
    from ..excel.analyzer import analyze_file
    result = analyze_file(sf.stored_path, str(settings.storage_dir / "work"), role_over or None, map_over or None)
    if not result["estimate_sheets"]:
        raise HTTPException(400, "계약내역으로 사용할 공종별 내역서 시트를 찾지 못했습니다. 시트 역할을 지정하세요.")
    try:
        cv = import_contract(db, p, sf, result, body.label, body.include_sheets, body.version_no)
    except ValueError as e:
        raise HTTPException(400, str(e))
    audit(db, user, "contract.import", p.id, target=f"v{cv.version_no} {cv.label}", after=cv.totals)
    db.commit()
    return {"version_id": cv.id, "version_no": cv.version_no, "label": cv.label, "totals": cv.totals,
            "item_count": sum(1 for i in cv.items if i.row_type == "item")}


@router.get("/mapping-rules")
def mapping_rules(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return [{"id": r.id, "name": r.name, "role": r.role, "signature": r.signature, "column_map": r.column_map}
            for r in db.query(MappingRule).order_by(MappingRule.id.desc())]
